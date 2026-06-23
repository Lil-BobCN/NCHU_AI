from __future__ import annotations

import re
import sys
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Pt


ROOT = Path(__file__).resolve().parents[1]
DOCS_DIR = ROOT / "docs"
API_MD = DOCS_DIR / "接口测试文档-2026-06-12.md"
API_DOCX = DOCS_DIR / "接口测试文档-2026-06-12.docx"
DB_MD = DOCS_DIR / "数据库表结构文档-2026-06-12.md"
DB_DOCX = DOCS_DIR / "数据库表结构文档-2026-06-12.docx"

BASE_URL = "http://127.0.0.1:8010"
TOKEN = "$TOKEN"
DOC_ID = "$DOCUMENT_ID"
JOB_ID = "$JOB_ID"
QA_ID = "$QA_PAIR_ID"
CONV_ID = "$CONVERSATION_ID"
LOG_ID = "$LOG_ID"
RUN_ID = "$RUN_ID"


@dataclass
class Field:
    name: str
    type: str
    required: str
    desc: str


@dataclass
class Endpoint:
    title: str
    method: str
    path: str
    auth: str
    desc: str
    params: list[Field]
    request: list[Field]
    response: list[Field]
    curl: str
    notes: list[str] | None = None


def md_table(headers: list[str], rows: list[list[str]]) -> str:
    if not rows:
        return "无\n"
    out = [
        "| " + " | ".join(headers) + " |",
        "| " + " | ".join(["---"] * len(headers)) + " |",
    ]
    for row in rows:
        out.append("| " + " | ".join(cell.replace("\n", "<br>") for cell in row) + " |")
    return "\n".join(out) + "\n"


def f(name: str, type_: str, required: str, desc: str) -> Field:
    return Field(name, type_, required, desc)


COMMON_OK_FIELDS = [
    f("code", "integer", "是", "统一业务状态码，成功固定为 0。"),
    f("message", "string", "是", "统一响应消息，成功默认 ok。"),
    f("data", "any", "否", "接口实际数据；不同接口结构不同。"),
]

PAGING_FIELDS = [
    f("items", "array", "是", "当前页数据列表。"),
    f("page", "integer", "是", "当前页码。"),
    f("page_size", "integer", "是", "每页数量。"),
    f("total", "integer", "是", "符合条件的总记录数。"),
]

DOCUMENT_FIELDS = [
    f("id", "uuid", "是", "文档 ID。"),
    f("title", "string", "是", "展示标题；标题乱码时后端会回退为文件名。"),
    f("file_name", "string", "是", "原始文件名。"),
    f("file_ext", "string", "是", "文件扩展名，小写。"),
    f("file_size", "integer", "是", "文件大小，单位字节。"),
    f("status", "string", "是", "文档状态：uploaded、parsed、chunked、indexed、failed、needs_conversion、needs_extraction、no_indexable_content、converted、extracted、deleted。"),
    f("source_url", "string|null", "否", "外部来源 URL。"),
    f("preview_url", "string|null", "否", "可预览访问地址，通常为 /api/v1/files 代理地址。"),
    f("download_url", "string|null", "否", "下载访问地址。"),
    f("parse_quality_score", "number", "是", "解析质量分，0 表示无质量分。"),
    f("error_message", "string|null", "否", "失败原因。"),
    f("created_at", "datetime|null", "否", "创建时间，ISO 8601。"),
    f("updated_at", "datetime|null", "否", "更新时间，ISO 8601。"),
]

JOB_FIELDS = [
    f("id", "uuid", "是", "任务 ID。"),
    f("document_id", "uuid", "否", "关联文档 ID。"),
    f("job_type", "string", "是", "任务类型：parse、chunk、embed、convert_office、extract_import、qa_generate。"),
    f("status", "string", "是", "任务状态：pending、running、succeeded、failed、skipped、canceled。"),
    f("progress", "integer", "是", "进度百分比，0-100。"),
    f("message", "string|null", "否", "任务阶段说明。"),
    f("error_message", "string|null", "否", "失败信息。"),
    f("result", "object", "是", "任务结果扩展字段，如转换后的文档 ID、解压导入文档列表等。"),
    f("created_at", "datetime|null", "否", "任务创建时间。"),
]

QA_FIELDS = [
    f("id", "uuid", "是", "QA ID。"),
    f("question", "string", "是", "问题。"),
    f("answer", "string", "是", "答案。"),
    f("status", "string", "是", "QA 状态，常用 enabled、disabled、draft。"),
    f("source_document_id", "uuid|null", "否", "来源文档 ID。"),
    f("source_chunk_ids", "uuid[]", "否", "来源切片 ID 列表。"),
    f("source_url", "string|null", "否", "来源 URL。"),
    f("tags", "string[]", "否", "标签。"),
    f("version", "integer", "是", "版本号，更新时递增。"),
    f("updated_at", "datetime|null", "否", "更新时间。"),
]

CONV_FIELDS = [
    f("id", "uuid", "是", "会话 ID。"),
    f("title", "string", "是", "会话标题。"),
    f("summary", "string", "是", "会话摘要。"),
    f("message_count", "integer", "是", "消息数量。"),
    f("last_message_at", "datetime|null", "否", "最后消息时间。"),
    f("created_at", "datetime|null", "否", "创建时间。"),
]

MESSAGE_FIELDS = [
    f("id", "uuid", "是", "消息 ID。"),
    f("conversation_id", "uuid", "是", "所属会话 ID。"),
    f("role", "string", "是", "消息角色：user 或 assistant。"),
    f("content", "string", "是", "消息正文。"),
    f("rewritten_query", "string|null", "否", "用于检索的改写查询。"),
    f("retrieval_trace", "object", "是", "检索过程追踪。"),
    f("citations", "array", "是", "引用来源列表。"),
    f("suggested_questions", "array", "是", "推荐追问列表。"),
    f("created_at", "datetime|null", "否", "创建时间。"),
]


def endpoints() -> list[Endpoint]:
    return [
        Endpoint(
            "根路径",
            "GET",
            "/",
            "否",
            "服务根路径，用于快速确认应用名称和 API 进程是否响应。",
            [],
            [],
            [f("name", "string", "是", "应用名称。"), f("status", "string", "是", "固定 ok。")],
            f"curl -s {BASE_URL}/",
        ),
        Endpoint(
            "健康检查",
            "GET",
            "/api/v1/health",
            "否",
            "检查 API、PostgreSQL、Redis、MinIO 是否可用。",
            [],
            [],
            COMMON_OK_FIELDS + [
                f("data.api", "string", "是", "API 状态。"),
                f("data.postgres", "string", "是", "PostgreSQL 状态，ok 或 error 信息。"),
                f("data.redis", "string", "是", "Redis 状态，ok 或 error 信息。"),
                f("data.minio", "string", "是", "MinIO 状态，ok 或 error 信息。"),
            ],
            f"curl -s {BASE_URL}/api/v1/health",
        ),
        Endpoint(
            "就绪检查",
            "GET",
            "/api/v1/readiness",
            "否",
            "与健康检查逻辑相同，可用于部署探针或测试环境依赖验证。",
            [],
            [],
            COMMON_OK_FIELDS + [
                f("data.api", "string", "是", "API 状态。"),
                f("data.postgres", "string", "是", "PostgreSQL 状态。"),
                f("data.redis", "string", "是", "Redis 状态。"),
                f("data.minio", "string", "是", "MinIO 状态。"),
            ],
            f"curl -s {BASE_URL}/api/v1/readiness",
        ),
        Endpoint(
            "文件代理下载/预览",
            "GET",
            "/api/v1/files/{bucket}/{object_key}",
            "否",
            "从 MinIO 读取对象并通过后端代理返回文件流。对象 key 支持路径。",
            [
                f("bucket", "path string", "是", "MinIO bucket 名，如 rag-documents。"),
                f("object_key", "path string", "是", "对象 key，需 URL 编码。"),
                f("download_name", "query string", "否", "传入时响应 Content-Disposition 下载文件名。"),
            ],
            [],
            [
                f("Content-Type", "header", "是", "对象 MIME 类型，默认 application/octet-stream。"),
                f("Content-Length", "header", "否", "对象字节数。"),
                f("Content-Disposition", "header", "否", "download_name 存在时返回附件下载头。"),
                f("body", "binary stream", "是", "文件内容。"),
            ],
            f"curl -L \"{BASE_URL}/api/v1/files/rag-documents/2026/06/demo.pdf?download_name=demo.pdf\" -o demo.pdf",
            ["该接口不走统一 code/message/data 包装。", "文件不存在或无权限时返回 HTTP 404，detail=file_not_found_or_inaccessible。"],
        ),
        Endpoint(
            "文件代理 HEAD",
            "HEAD",
            "/api/v1/files/{bucket}/{object_key}",
            "否",
            "只检查文件元信息，不下载文件体。",
            [
                f("bucket", "path string", "是", "MinIO bucket 名。"),
                f("object_key", "path string", "是", "对象 key。"),
                f("download_name", "query string", "否", "下载文件名。"),
            ],
            [],
            [
                f("HTTP 200", "status", "是", "对象可访问。"),
                f("Content-Type", "header", "是", "对象 MIME 类型。"),
                f("Content-Length", "header", "否", "对象字节数。"),
            ],
            f"curl -I \"{BASE_URL}/api/v1/files/rag-documents/2026/06/demo.pdf\"",
            ["该接口不返回响应体。"],
        ),
        Endpoint(
            "登录",
            "POST",
            "/api/v1/auth/login",
            "否",
            "管理员登录，成功返回 JWT token。后续受保护接口使用 Authorization: Bearer <token>。",
            [],
            [
                f("username", "string", "是", "管理员账号，默认配置常见为 admin。"),
                f("password", "string", "是", "管理员密码。"),
            ],
            COMMON_OK_FIELDS + [
                f("data.access_token", "string", "是", "JWT 访问令牌。"),
                f("data.token_type", "string", "是", "固定 bearer。"),
                f("data.expires_in", "integer", "是", "有效期秒数，当前为 86400。"),
                f("data.admin.id", "uuid", "是", "管理员 ID。"),
                f("data.admin.username", "string", "是", "用户名。"),
                f("data.admin.display_name", "string", "是", "显示名。"),
            ],
            f"""curl -s -X POST {BASE_URL}/api/v1/auth/login \\
  -H "Content-Type: application/json" \\
  -d '{{"username":"admin","password":"admin123"}}'""",
            ["账号或密码错误返回 HTTP 401。"],
        ),
        Endpoint(
            "当前登录用户",
            "GET",
            "/api/v1/auth/me",
            "是",
            "获取当前 token 对应管理员信息。",
            [],
            [],
            COMMON_OK_FIELDS + [
                f("data.id", "uuid", "是", "管理员 ID。"),
                f("data.username", "string", "是", "用户名。"),
                f("data.display_name", "string", "是", "显示名。"),
            ],
            f"""curl -s {BASE_URL}/api/v1/auth/me \\
  -H "Authorization: Bearer {TOKEN}\"""",
        ),
        Endpoint(
            "上传文档",
            "POST",
            "/api/v1/documents/upload",
            "是",
            "上传文档到 MinIO 并创建文档记录。auto_process=true 时创建后台解析/切片/向量化流水线。",
            [],
            [
                f("file", "multipart file", "是", "上传文件，最大 100MB。支持 pdf、doc、docx、xls、xlsx、xlsm、txt、md、jpg、jpeg、png、bmp、tif、tiff、zip、rar。"),
                f("title", "form string", "否", "文档标题，缺省使用文件名。"),
                f("source_url", "form string", "否", "外部来源 URL。"),
                f("auto_process", "form boolean", "否", "是否自动处理，默认 true。"),
            ],
            COMMON_OK_FIELDS + [
                f("data.document", "object", "是", "文档对象，字段见通用 Document。"),
                f("data.jobs", "array", "是", "自动处理任务概要；auto_process=false 时为空数组。"),
            ] + DOCUMENT_FIELDS,
            f"""curl -s -X POST {BASE_URL}/api/v1/documents/upload \\
  -H "Authorization: Bearer {TOKEN}" \\
  -F "file=@sample_contract.pdf" \\
  -F "title=测试合同" \\
  -F "auto_process=true\"""",
            ["不支持文件类型返回 HTTP 415。", "超过 max_upload_size_mb 返回 HTTP 413。"],
        ),
        Endpoint(
            "文档列表",
            "GET",
            "/api/v1/documents",
            "是",
            "分页查询未删除文档。",
            [
                f("keyword", "query string", "否", "按 title 或 file_name 模糊搜索。"),
                f("status", "query string", "否", "按文档状态过滤。"),
                f("page", "query integer", "否", "页码，默认 1。"),
                f("page_size", "query integer", "否", "每页数量，默认 20。"),
            ],
            [],
            COMMON_OK_FIELDS + PAGING_FIELDS + DOCUMENT_FIELDS,
            f"""curl -s "{BASE_URL}/api/v1/documents?keyword=合同&status=indexed&page=1&page_size=10" \\
  -H "Authorization: Bearer {TOKEN}\"""",
        ),
        Endpoint(
            "文档详情",
            "GET",
            "/api/v1/documents/{document_id}",
            "是",
            "获取单个文档详情、切片数量和最近 5 条任务。",
            [f("document_id", "path uuid", "是", "文档 ID。")],
            [],
            COMMON_OK_FIELDS + DOCUMENT_FIELDS + [
                f("data.stats.chunk_count", "integer", "是", "文档切片数量。"),
                f("data.latest_jobs", "array", "是", "最近 5 条任务，字段见 Job。"),
            ],
            f"""curl -s {BASE_URL}/api/v1/documents/{DOC_ID} \\
  -H "Authorization: Bearer {TOKEN}\"""",
            ["document_id 格式错误返回 HTTP 422；文档不存在返回 HTTP 404。"],
        ),
        Endpoint(
            "删除文档",
            "DELETE",
            "/api/v1/documents/{document_id}",
            "是",
            "软删除文档并尝试删除 MinIO 原始文件和解析结果文件。",
            [f("document_id", "path uuid", "是", "文档 ID。")],
            [],
            COMMON_OK_FIELDS + [f("data.id", "uuid", "是", "文档 ID。"), f("data.status", "string", "是", "固定 deleted。")],
            f"""curl -s -X DELETE {BASE_URL}/api/v1/documents/{DOC_ID} \\
  -H "Authorization: Bearer {TOKEN}\"""",
        ),
        Endpoint(
            "重新解析文档",
            "POST",
            "/api/v1/documents/{document_id}/reparse",
            "是",
            "读取原始文件并重新执行完整流水线：解析、切片、向量化。",
            [f("document_id", "path uuid", "是", "文档 ID。")],
            [],
            COMMON_OK_FIELDS + [f("data.document_id", "uuid", "是", "文档 ID。"), f("data.status", "string", "是", "pending。")],
            f"""curl -s -X POST {BASE_URL}/api/v1/documents/{DOC_ID}/reparse \\
  -H "Authorization: Bearer {TOKEN}\"""",
        ),
        Endpoint(
            "重新切片并向量化",
            "POST",
            "/api/v1/documents/{document_id}/rechunk",
            "是",
            "基于已有解析结果重新切片，随后自动执行向量化。",
            [f("document_id", "path uuid", "是", "文档 ID。")],
            [],
            COMMON_OK_FIELDS + [
                f("data.document_id", "uuid", "是", "文档 ID。"),
                f("data.status", "string", "是", "pending 或 skipped。"),
                f("data.message", "string", "否", "skipped 时返回跳过原因。"),
            ],
            f"""curl -s -X POST {BASE_URL}/api/v1/documents/{DOC_ID}/rechunk \\
  -H "Authorization: Bearer {TOKEN}\"""",
        ),
        Endpoint(
            "重新向量化",
            "POST",
            "/api/v1/documents/{document_id}/reembed",
            "是",
            "清理当前文档 embedding 后重新生成向量。",
            [f("document_id", "path uuid", "是", "文档 ID。")],
            [],
            COMMON_OK_FIELDS + [
                f("data.document_id", "uuid", "是", "文档 ID。"),
                f("data.status", "string", "是", "pending 或 skipped。"),
                f("data.message", "string", "否", "skipped 时返回跳过原因。"),
            ],
            f"""curl -s -X POST {BASE_URL}/api/v1/documents/{DOC_ID}/reembed \\
  -H "Authorization: Bearer {TOKEN}\"""",
        ),
        Endpoint(
            "转换旧版 Office 文档",
            "POST",
            "/api/v1/documents/{document_id}/convert-office",
            "是",
            "将 .doc/.xls 转为新版 Office 文件并创建子文档，随后对子文档入库。",
            [f("document_id", "path uuid", "是", "文档 ID，文件扩展名必须为 .doc 或 .xls。")],
            [],
            COMMON_OK_FIELDS + [
                f("data.document_id", "uuid", "是", "源文档 ID。"),
                f("data.job_type", "string", "是", "convert_office。"),
                f("data.status", "string", "是", "pending。"),
            ],
            f"""curl -s -X POST {BASE_URL}/api/v1/documents/{DOC_ID}/convert-office \\
  -H "Authorization: Bearer {TOKEN}\"""",
            ["非 .doc/.xls 返回 HTTP 400。"],
        ),
        Endpoint(
            "解压导入压缩包",
            "POST",
            "/api/v1/documents/{document_id}/extract-archive",
            "是",
            "安全解压 .zip/.rar 中支持的内部文件，为每个内部文件创建新文档并继续处理。",
            [f("document_id", "path uuid", "是", "文档 ID，文件扩展名必须为 .zip 或 .rar。")],
            [],
            COMMON_OK_FIELDS + [
                f("data.document_id", "uuid", "是", "压缩包文档 ID。"),
                f("data.job_type", "string", "是", "extract_import。"),
                f("data.status", "string", "是", "pending。"),
            ],
            f"""curl -s -X POST {BASE_URL}/api/v1/documents/{DOC_ID}/extract-archive \\
  -H "Authorization: Bearer {TOKEN}\"""",
            ["非 .zip/.rar 返回 HTTP 400。"],
        ),
        Endpoint(
            "从文档生成 QA",
            "POST",
            "/api/v1/documents/{document_id}/qa-pairs/generate",
            "是",
            "基于文档切片后台生成 QA 对，并同步 QA embedding。",
            [
                f("document_id", "path uuid", "是", "文档 ID。"),
                f("count", "query integer", "否", "生成数量，默认 10。"),
                f("auto_enable", "query boolean", "否", "是否直接启用，默认 true；false 时状态为 draft。"),
            ],
            [],
            COMMON_OK_FIELDS + [
                f("data.document_id", "uuid", "是", "文档 ID。"),
                f("data.job_type", "string", "是", "qa_generate。"),
                f("data.status", "string", "是", "pending 或 skipped。"),
            ],
            f"""curl -s -X POST "{BASE_URL}/api/v1/documents/{DOC_ID}/qa-pairs/generate?count=5&auto_enable=true" \\
  -H "Authorization: Bearer {TOKEN}\"""",
        ),
        Endpoint(
            "获取解析结果",
            "GET",
            "/api/v1/documents/{document_id}/parse-result",
            "是",
            "获取文档解析后的正文或 Markdown 内容。",
            [
                f("document_id", "path uuid", "是", "文档 ID。"),
                f("format", "query string", "否", "markdown 或 text，默认 markdown。"),
            ],
            [],
            COMMON_OK_FIELDS + [
                f("data.document_id", "uuid", "是", "文档 ID。"),
                f("data.format", "string", "是", "请求的格式。"),
                f("data.content", "string", "是", "解析内容。"),
                f("data.quality_score", "number", "是", "解析质量分。"),
                f("data.parse_meta", "object", "是", "解析元数据。"),
            ],
            f"""curl -s "{BASE_URL}/api/v1/documents/{DOC_ID}/parse-result?format=markdown" \\
  -H "Authorization: Bearer {TOKEN}\"""",
            ["解析结果不存在返回 HTTP 404。"],
        ),
        Endpoint(
            "文档切片列表",
            "GET",
            "/api/v1/documents/{document_id}/chunks",
            "是",
            "分页查询某个文档的切片。",
            [
                f("document_id", "path uuid", "是", "文档 ID。"),
                f("page", "query integer", "否", "页码，默认 1。"),
                f("page_size", "query integer", "否", "每页数量，默认 20。"),
                f("keyword", "query string", "否", "按切片 content 模糊搜索。"),
            ],
            [],
            COMMON_OK_FIELDS + PAGING_FIELDS + [
                f("items[].id", "uuid", "是", "切片 ID。"),
                f("items[].chunk_no", "integer", "是", "切片序号。"),
                f("items[].chunk_type", "string", "是", "切片类型，如 text、table、parent。"),
                f("items[].content", "string", "是", "切片内容。"),
                f("items[].char_count", "integer", "是", "字符数。"),
                f("items[].token_count", "integer|null", "否", "token 数。"),
                f("items[].page_start", "integer|null", "否", "起始页码。"),
                f("items[].page_end", "integer|null", "否", "结束页码。"),
                f("items[].section_path", "string|null", "否", "章节路径。"),
                f("items[].metadata", "object", "是", "切片元数据。"),
                f("items[].is_active", "boolean", "是", "是否参与检索。"),
            ],
            f"""curl -s "{BASE_URL}/api/v1/documents/{DOC_ID}/chunks?page=1&page_size=20&keyword=付款" \\
  -H "Authorization: Bearer {TOKEN}\"""",
        ),
        Endpoint(
            "任务详情",
            "GET",
            "/api/v1/jobs/{job_id}",
            "是",
            "查询单个后台任务进度。优先从 Redis 读取缓存进度，缓存不存在时读取数据库。",
            [f("job_id", "path uuid", "是", "任务 ID。")],
            [],
            COMMON_OK_FIELDS + JOB_FIELDS,
            f"""curl -s {BASE_URL}/api/v1/jobs/{JOB_ID} \\
  -H "Authorization: Bearer {TOKEN}\"""",
        ),
        Endpoint(
            "文档任务列表",
            "GET",
            "/api/v1/documents/{document_id}/jobs",
            "是",
            "查询某个文档的所有后台任务，按创建时间倒序。",
            [f("document_id", "path uuid", "是", "文档 ID。")],
            [],
            COMMON_OK_FIELDS + JOB_FIELDS,
            f"""curl -s {BASE_URL}/api/v1/documents/{DOC_ID}/jobs \\
  -H "Authorization: Bearer {TOKEN}\"""",
        ),
        Endpoint(
            "QA 列表",
            "GET",
            "/api/v1/qa-pairs",
            "是",
            "分页查询 QA 对。",
            [
                f("keyword", "query string", "否", "按 question 或 answer 模糊搜索。"),
                f("status", "query string", "否", "按 QA 状态过滤。"),
                f("document_id", "query uuid", "否", "按来源文档过滤。"),
                f("page", "query integer", "否", "页码，默认 1。"),
                f("page_size", "query integer", "否", "每页数量，默认 20。"),
            ],
            [],
            COMMON_OK_FIELDS + PAGING_FIELDS + QA_FIELDS,
            f"""curl -s "{BASE_URL}/api/v1/qa-pairs?keyword=报名&status=enabled&page=1&page_size=10" \\
  -H "Authorization: Bearer {TOKEN}\"""",
        ),
        Endpoint(
            "创建 QA",
            "POST",
            "/api/v1/qa-pairs",
            "是",
            "手动创建 QA 对，并在后台同步 QA embedding。",
            [],
            [
                f("question", "string", "是", "问题。"),
                f("answer", "string", "是", "答案。"),
                f("status", "string", "否", "默认 enabled。"),
                f("source_document_id", "uuid|null", "否", "来源文档 ID。"),
                f("source_chunk_ids", "uuid[]|null", "否", "来源切片 ID。"),
                f("tags", "string[]|null", "否", "标签。"),
            ],
            COMMON_OK_FIELDS + QA_FIELDS,
            f"""curl -s -X POST {BASE_URL}/api/v1/qa-pairs \\
  -H "Authorization: Bearer {TOKEN}" \\
  -H "Content-Type: application/json" \\
  -d '{{"question":"报名时间是什么时候？","answer":"以通知文件为准。","status":"enabled","tags":["测试"]}}'""",
        ),
        Endpoint(
            "更新 QA",
            "PUT",
            "/api/v1/qa-pairs/{qa_pair_id}",
            "是",
            "更新 QA 内容、状态或标签，版本号递增，并后台同步 embedding。",
            [f("qa_pair_id", "path uuid", "是", "QA ID。")],
            [
                f("question", "string|null", "否", "新问题。"),
                f("answer", "string|null", "否", "新答案。"),
                f("status", "string|null", "否", "新状态。"),
                f("tags", "string[]|null", "否", "新标签。"),
            ],
            COMMON_OK_FIELDS + QA_FIELDS,
            f"""curl -s -X PUT {BASE_URL}/api/v1/qa-pairs/{QA_ID} \\
  -H "Authorization: Bearer {TOKEN}" \\
  -H "Content-Type: application/json" \\
  -d '{{"answer":"更新后的答案","tags":["回归测试"]}}'""",
        ),
        Endpoint(
            "删除 QA",
            "DELETE",
            "/api/v1/qa-pairs/{qa_pair_id}",
            "是",
            "软删除 QA：设置 deleted_at，并将 status 改为 disabled。",
            [f("qa_pair_id", "path uuid", "是", "QA ID。")],
            [],
            COMMON_OK_FIELDS + [f("data.id", "uuid", "是", "QA ID。")],
            f"""curl -s -X DELETE {BASE_URL}/api/v1/qa-pairs/{QA_ID} \\
  -H "Authorization: Bearer {TOKEN}\"""",
        ),
        Endpoint(
            "更新 QA 状态",
            "PATCH",
            "/api/v1/qa-pairs/{qa_pair_id}/status",
            "是",
            "只更新 QA 状态。",
            [f("qa_pair_id", "path uuid", "是", "QA ID。")],
            [f("status", "string", "是", "目标状态，如 enabled、disabled、draft。")],
            COMMON_OK_FIELDS + QA_FIELDS,
            f"""curl -s -X PATCH {BASE_URL}/api/v1/qa-pairs/{QA_ID}/status \\
  -H "Authorization: Bearer {TOKEN}" \\
  -H "Content-Type: application/json" \\
  -d '{{"status":"disabled"}}'""",
        ),
        Endpoint(
            "创建会话",
            "POST",
            "/api/v1/conversations",
            "是",
            "创建一个聊天会话。",
            [],
            [f("title", "string", "否", "会话标题，默认“新的对话”。")],
            COMMON_OK_FIELDS + CONV_FIELDS,
            f"""curl -s -X POST {BASE_URL}/api/v1/conversations \\
  -H "Authorization: Bearer {TOKEN}" \\
  -H "Content-Type: application/json" \\
  -d '{{"title":"测试会话"}}'""",
        ),
        Endpoint(
            "会话列表",
            "GET",
            "/api/v1/conversations",
            "是",
            "分页查询未删除会话。",
            [f("page", "query integer", "否", "页码，默认 1。"), f("page_size", "query integer", "否", "每页数量，默认 20。")],
            [],
            COMMON_OK_FIELDS + PAGING_FIELDS + CONV_FIELDS,
            f"""curl -s "{BASE_URL}/api/v1/conversations?page=1&page_size=20" \\
  -H "Authorization: Bearer {TOKEN}\"""",
        ),
        Endpoint(
            "会话消息列表",
            "GET",
            "/api/v1/conversations/{conversation_id}/messages",
            "是",
            "按时间正序查询某个会话的消息。",
            [f("conversation_id", "path uuid", "是", "会话 ID。")],
            [],
            COMMON_OK_FIELDS + MESSAGE_FIELDS,
            f"""curl -s {BASE_URL}/api/v1/conversations/{CONV_ID}/messages \\
  -H "Authorization: Bearer {TOKEN}\"""",
        ),
        Endpoint(
            "删除会话",
            "DELETE",
            "/api/v1/conversations/{conversation_id}",
            "是",
            "软删除会话。",
            [f("conversation_id", "path uuid", "是", "会话 ID。")],
            [],
            COMMON_OK_FIELDS + [f("data.id", "uuid", "是", "会话 ID。")],
            f"""curl -s -X DELETE {BASE_URL}/api/v1/conversations/{CONV_ID} \\
  -H "Authorization: Bearer {TOKEN}\"""",
        ),
        Endpoint(
            "检索调试搜索",
            "POST",
            "/api/v1/retrieval/search",
            "是",
            "执行混合检索：向量召回、关键词召回、QA 召回、RRF 融合、rerank 和上下文扩展。",
            [],
            [
                f("query", "string", "是", "检索问题。"),
                f("top_k", "integer", "否", "召回数量，默认 30，最大 100。"),
                f("rerank_top_k", "integer", "否", "最终返回数量，默认 5，最大 50。"),
                f("document_ids", "uuid[]|null", "否", "限定文档 ID。"),
                f("enable_rewrite", "boolean", "否", "当前接口定义包含该字段，但 search 实现未使用。"),
                f("enable_qa_recall", "boolean", "否", "是否启用 QA 召回，默认 true。"),
                f("enable_keyword_recall", "boolean", "否", "是否启用关键词召回，默认 true。"),
                f("enable_vector_recall", "boolean", "否", "是否启用向量召回，默认 true。"),
            ],
            COMMON_OK_FIELDS + [
                f("data.raw_query", "string", "是", "原始查询。"),
                f("data.rewritten_query", "string", "是", "当前同 raw_query。"),
                f("data.recall_results", "array", "是", "召回结果。"),
                f("data.rerank_results", "array", "是", "rerank 后结果。"),
                f("data.final_context", "array", "是", "扩展后的最终上下文。"),
                f("data.answer_context", "array", "是", "用于回答的精选上下文。"),
                f("data.citations", "array", "是", "引用来源。"),
                f("data.retrieval_options", "object", "是", "本次检索参数。"),
                f("data.effective_settings", "object", "是", "生效的 RAG 设置。"),
            ],
            f"""curl -s -X POST {BASE_URL}/api/v1/retrieval/search \\
  -H "Authorization: Bearer {TOKEN}" \\
  -H "Content-Type: application/json" \\
  -d '{{"query":"报名与缴费是怎样的？","top_k":30,"rerank_top_k":5}}'""",
        ),
        Endpoint(
            "检索日志列表",
            "GET",
            "/api/v1/retrieval/logs",
            "是",
            "分页查询检索日志概要。",
            [
                f("keyword", "query string", "否", "按 raw_query 模糊搜索。"),
                f("page", "query integer", "否", "页码，默认 1。"),
                f("page_size", "query integer", "否", "每页数量，默认 20。"),
            ],
            [],
            COMMON_OK_FIELDS + PAGING_FIELDS + [
                f("items[].id", "uuid", "是", "日志 ID。"),
                f("items[].raw_query", "string", "是", "原始问题。"),
                f("items[].rewritten_query", "string|null", "否", "改写后问题。"),
                f("items[].latency_ms", "integer|null", "否", "耗时毫秒。"),
                f("items[].created_at", "datetime|null", "否", "创建时间。"),
            ],
            f"""curl -s "{BASE_URL}/api/v1/retrieval/logs?keyword=报名&page=1&page_size=20" \\
  -H "Authorization: Bearer {TOKEN}\"""",
        ),
        Endpoint(
            "检索日志详情",
            "GET",
            "/api/v1/retrieval/logs/{log_id}",
            "是",
            "查询单条检索日志完整追踪。不存在时 code=0 且 data=null。",
            [f("log_id", "path uuid", "是", "日志 ID。")],
            [],
            COMMON_OK_FIELDS + [
                f("data.id", "uuid", "是", "日志 ID。"),
                f("data.raw_query", "string", "是", "原始问题。"),
                f("data.rewritten_query", "string|null", "否", "改写问题。"),
                f("data.recall_results", "array", "是", "召回结果。"),
                f("data.rerank_results", "array", "是", "排序结果。"),
                f("data.final_context", "array", "是", "最终上下文。"),
                f("data.citations", "array", "是", "引用。"),
                f("data.suggested_questions", "array", "是", "推荐追问。"),
                f("data.answer", "string|null", "否", "回答文本。"),
                f("data.latency_ms", "integer|null", "否", "耗时毫秒。"),
            ],
            f"""curl -s {BASE_URL}/api/v1/retrieval/logs/{LOG_ID} \\
  -H "Authorization: Bearer {TOKEN}\"""",
        ),
        Endpoint(
            "流式问答",
            "POST",
            "/api/v1/chat/stream",
            "是",
            "SSE 流式问答。会创建或复用会话，写入用户消息、助手消息和检索日志。",
            [],
            [
                f("conversation_id", "uuid|null", "否", "会话 ID；为空时自动创建。"),
                f("question", "string", "是", "用户问题。"),
                f("top_k", "integer", "否", "召回数量，默认 8。"),
                f("rerank_top_k", "integer", "否", "最终上下文数量，默认 5。"),
                f("document_ids", "uuid[]|null", "否", "限定文档。"),
                f("enable_rewrite", "boolean", "否", "是否允许结合历史改写追问，默认 true。"),
                f("enable_suggested_questions", "boolean", "否", "是否生成推荐追问，默认 true。"),
            ],
            [
                f("event: message_start", "SSE", "是", "返回 conversation_id、user_message_id、assistant_message_id 等。"),
                f("event: retrieval_start", "SSE", "是", "检索开始。"),
                f("event: retrieval", "SSE", "是", "召回和排序数量。"),
                f("event: retrieval_done", "SSE", "是", "检索完成。"),
                f("event: answer_cache", "SSE", "否", "命中答案缓存时返回。"),
                f("event: delta", "SSE", "是", "回答增量文本，data.content。"),
                f("event: citations", "SSE", "是", "引用来源。"),
                f("event: suggested_questions", "SSE", "是", "推荐追问。"),
                f("event: message_end", "SSE", "是", "回答结束，含 latency_ms。"),
            ],
            f"""curl -N -X POST {BASE_URL}/api/v1/chat/stream \\
  -H "Authorization: Bearer {TOKEN}" \\
  -H "Content-Type: application/json" \\
  -d '{{"question":"辅修学士学位报名与缴费是怎样的？","enable_suggested_questions":false}}'""",
            ["该接口 media_type 为 text/event-stream，不使用统一响应包装。"],
        ),
        Endpoint(
            "非流式问答",
            "POST",
            "/api/v1/chat",
            "是",
            "一次性返回完整回答、引用、推荐追问和检索追踪。",
            [],
            [
                f("conversation_id", "uuid|null", "否", "会话 ID；为空时自动创建。"),
                f("question", "string", "是", "用户问题。"),
                f("top_k", "integer", "否", "召回数量，默认 8。"),
                f("rerank_top_k", "integer", "否", "最终上下文数量，默认 5。"),
                f("document_ids", "uuid[]|null", "否", "限定文档。"),
                f("enable_rewrite", "boolean", "否", "是否允许历史改写，默认 true。"),
                f("enable_suggested_questions", "boolean", "否", "是否生成推荐追问，默认 true。"),
            ],
            COMMON_OK_FIELDS + [
                f("data.answer", "string", "是", "完整回答。"),
                f("data.citations", "array", "是", "引用来源。"),
                f("data.suggested_questions", "array", "是", "推荐追问。"),
                f("data.retrieval_trace", "object", "是", "检索追踪。"),
            ],
            f"""curl -s -X POST {BASE_URL}/api/v1/chat \\
  -H "Authorization: Bearer {TOKEN}" \\
  -H "Content-Type: application/json" \\
  -d '{{"question":"辅修学士学位报名与缴费是怎样的？","top_k":8,"rerank_top_k":5}}'""",
        ),
        Endpoint(
            "获取 RAG 设置",
            "GET",
            "/api/v1/settings/rag",
            "是",
            "获取当前生效的 RAG 参数。",
            [],
            [],
            COMMON_OK_FIELDS + [
                f("data.chunk_size", "integer", "是", "普通文本切片长度。"),
                f("data.chunk_overlap", "integer", "是", "切片重叠长度。"),
                f("data.table_chunk_rows", "integer", "是", "表格分块行数。"),
                f("data.table_full_index_max_rows", "integer", "是", "表格全量索引最大行数。"),
                f("data.table_sample_rows", "integer", "是", "表格抽样行数。"),
                f("data.max_document_chunks", "integer", "是", "单文档最大切片数。"),
                f("data.max_embedding_chars_per_document", "integer", "是", "单文档最大向量化字符数。"),
                f("data.vector_top_k", "integer", "是", "向量召回数。"),
                f("data.keyword_top_k", "integer", "是", "关键词召回数。"),
                f("data.qa_top_k", "integer", "是", "QA 召回数。"),
                f("data.rerank_top_k", "integer", "是", "rerank 返回数。"),
                f("data.similarity_threshold", "number", "是", "相似度阈值。"),
                f("data.rerank_threshold", "number", "是", "rerank 阈值。"),
                f("data.rerank_enabled", "boolean", "是", "是否启用 rerank。"),
                f("data.rerank_max_candidates", "integer", "是", "rerank 最大候选数量。"),
                f("data.persisted", "boolean", "是", "固定 true。"),
            ],
            f"""curl -s {BASE_URL}/api/v1/settings/rag \\
  -H "Authorization: Bearer {TOKEN}\"""",
        ),
        Endpoint(
            "更新 RAG 设置",
            "PUT",
            "/api/v1/settings/rag",
            "是",
            "保存 RAG 参数到 .cache/rag_settings.json，并返回当前生效值。",
            [],
            [
                f("chunk_size", "integer", "是", "普通文本切片长度。"),
                f("chunk_overlap", "integer", "是", "切片重叠长度。"),
                f("table_chunk_rows", "integer", "否", "默认 50。"),
                f("table_full_index_max_rows", "integer", "否", "默认 500。"),
                f("table_sample_rows", "integer", "否", "默认 50。"),
                f("max_document_chunks", "integer", "否", "默认 300。"),
                f("max_embedding_chars_per_document", "integer", "否", "默认 200000。"),
                f("vector_top_k", "integer", "是", "向量召回数。"),
                f("keyword_top_k", "integer", "是", "关键词召回数。"),
                f("qa_top_k", "integer", "是", "QA 召回数。"),
                f("rerank_top_k", "integer", "是", "rerank 返回数。"),
                f("similarity_threshold", "number", "是", "相似度阈值。"),
                f("rerank_threshold", "number", "是", "rerank 阈值。"),
                f("rerank_enabled", "boolean", "否", "默认 true。"),
                f("rerank_max_candidates", "integer", "否", "默认 10。"),
            ],
            COMMON_OK_FIELDS + [f("data", "object", "是", "字段同 GET /settings/rag。")],
            f"""curl -s -X PUT {BASE_URL}/api/v1/settings/rag \\
  -H "Authorization: Bearer {TOKEN}" \\
  -H "Content-Type: application/json" \\
  -d '{{"chunk_size":800,"chunk_overlap":120,"table_chunk_rows":50,"table_full_index_max_rows":500,"table_sample_rows":50,"max_document_chunks":300,"max_embedding_chars_per_document":200000,"vector_top_k":30,"keyword_top_k":30,"qa_top_k":10,"rerank_top_k":5,"similarity_threshold":0.35,"rerank_threshold":0.45,"rerank_enabled":true,"rerank_max_candidates":10}}'""",
        ),
        Endpoint(
            "评测用例列表",
            "GET",
            "/api/v1/evaluation/cases",
            "是",
            "分页查询评测用例。",
            [f("page", "query integer", "否", "页码，默认 1。"), f("page_size", "query integer", "否", "每页数量，默认 20。")],
            [],
            COMMON_OK_FIELDS + PAGING_FIELDS + [
                f("items[].id", "uuid", "是", "用例 ID。"),
                f("items[].question", "string", "是", "测试问题。"),
                f("items[].expected_answer", "string|null", "否", "期望答案或关键词。"),
                f("items[].expected_document_ids", "uuid[]", "否", "期望命中文档。"),
                f("items[].expected_chunk_ids", "uuid[]", "否", "期望命中切片。"),
                f("items[].tags", "string[]", "否", "标签。"),
                f("items[].difficulty", "string|null", "否", "难度。"),
                f("items[].is_active", "boolean", "是", "是否启用。"),
                f("items[].created_at", "datetime|null", "否", "创建时间。"),
            ],
            f"""curl -s "{BASE_URL}/api/v1/evaluation/cases?page=1&page_size=20" \\
  -H "Authorization: Bearer {TOKEN}\"""",
        ),
        Endpoint(
            "创建评测用例",
            "POST",
            "/api/v1/evaluation/cases",
            "是",
            "创建 RAG 评测用例。",
            [],
            [
                f("question", "string", "是", "测试问题。"),
                f("expected_answer", "string|null", "否", "期望答案或关键词。"),
                f("expected_document_ids", "uuid[]|null", "否", "期望命中文档 ID。"),
                f("expected_chunk_ids", "uuid[]|null", "否", "期望命中切片 ID。"),
                f("tags", "string[]|null", "否", "标签。"),
                f("difficulty", "string|null", "否", "难度。"),
            ],
            COMMON_OK_FIELDS + [
                f("data.id", "uuid", "是", "用例 ID。"),
                f("data.question", "string", "是", "问题。"),
                f("data.expected_answer", "string|null", "否", "期望答案。"),
                f("data.expected_document_ids", "uuid[]", "否", "期望文档。"),
                f("data.expected_chunk_ids", "uuid[]", "否", "期望切片。"),
                f("data.tags", "string[]", "否", "标签。"),
                f("data.difficulty", "string|null", "否", "难度。"),
                f("data.is_active", "boolean", "是", "是否启用。"),
                f("data.created_at", "datetime|null", "否", "创建时间。"),
            ],
            f"""curl -s -X POST {BASE_URL}/api/v1/evaluation/cases \\
  -H "Authorization: Bearer {TOKEN}" \\
  -H "Content-Type: application/json" \\
  -d '{{"question":"报名与缴费是怎样的？","expected_answer":"报名,缴费","tags":["回归"],"difficulty":"normal"}}'""",
        ),
        Endpoint(
            "创建评测批次",
            "POST",
            "/api/v1/evaluation/runs",
            "是",
            "创建后台评测批次，异步执行用例并计算指标。",
            [],
            [
                f("name", "string", "是", "批次名称。"),
                f("case_ids", "uuid[]|null", "否", "指定用例 ID；为空则执行所有 active 用例。"),
                f("config", "object", "否", "评测配置，默认空对象。"),
            ],
            COMMON_OK_FIELDS + [
                f("data.run_id", "uuid", "是", "评测批次 ID。"),
                f("data.status", "string", "是", "running。"),
            ],
            f"""curl -s -X POST {BASE_URL}/api/v1/evaluation/runs \\
  -H "Authorization: Bearer {TOKEN}" \\
  -H "Content-Type: application/json" \\
  -d '{{"name":"回归评测-2026-06-12","case_ids":[],"config":{{"top_k":8}}}}'""",
        ),
        Endpoint(
            "评测批次详情",
            "GET",
            "/api/v1/evaluation/runs/{run_id}",
            "是",
            "查询评测批次状态、指标和结果列表。",
            [f("run_id", "path uuid", "是", "评测批次 ID。")],
            [],
            COMMON_OK_FIELDS + [
                f("data.id", "uuid", "是", "批次 ID。"),
                f("data.name", "string", "是", "批次名称。"),
                f("data.status", "string", "是", "running、succeeded 或 failed。"),
                f("data.metrics", "object", "是", "聚合指标，如 case_count、top3_hit_rate、answer_score_avg。"),
                f("data.results", "array", "是", "评测结果列表。"),
                f("results[].case_id", "uuid", "是", "用例 ID。"),
                f("results[].answer", "string|null", "否", "实际回答。"),
                f("results[].hit_top3", "boolean|null", "否", "Top3 是否命中期望来源。"),
                f("results[].hit_top5", "boolean|null", "否", "Top5 是否命中期望来源。"),
                f("results[].answer_score", "number|null", "否", "答案关键词得分。"),
                f("results[].citation_score", "number|null", "否", "引用得分。"),
                f("results[].suggested_question_score", "number|null", "否", "推荐追问得分。"),
                f("results[].failure_reason", "string|null", "否", "失败原因。"),
                f("results[].trace", "object", "是", "评测追踪。"),
            ],
            f"""curl -s {BASE_URL}/api/v1/evaluation/runs/{RUN_ID} \\
  -H "Authorization: Bearer {TOKEN}\"""",
        ),
    ]


EXPECTED_OPERATIONS = {
    ("GET", "/"),
    ("GET", "/api/v1/health"),
    ("GET", "/api/v1/readiness"),
    ("GET", "/api/v1/files/{bucket}/{object_key}"),
    ("HEAD", "/api/v1/files/{bucket}/{object_key}"),
    ("POST", "/api/v1/auth/login"),
    ("GET", "/api/v1/auth/me"),
    ("POST", "/api/v1/documents/upload"),
    ("GET", "/api/v1/documents"),
    ("GET", "/api/v1/documents/{document_id}"),
    ("DELETE", "/api/v1/documents/{document_id}"),
    ("POST", "/api/v1/documents/{document_id}/reparse"),
    ("POST", "/api/v1/documents/{document_id}/rechunk"),
    ("POST", "/api/v1/documents/{document_id}/reembed"),
    ("POST", "/api/v1/documents/{document_id}/convert-office"),
    ("POST", "/api/v1/documents/{document_id}/extract-archive"),
    ("POST", "/api/v1/documents/{document_id}/qa-pairs/generate"),
    ("GET", "/api/v1/documents/{document_id}/parse-result"),
    ("GET", "/api/v1/documents/{document_id}/chunks"),
    ("GET", "/api/v1/jobs/{job_id}"),
    ("GET", "/api/v1/documents/{document_id}/jobs"),
    ("GET", "/api/v1/qa-pairs"),
    ("POST", "/api/v1/qa-pairs"),
    ("PUT", "/api/v1/qa-pairs/{qa_pair_id}"),
    ("DELETE", "/api/v1/qa-pairs/{qa_pair_id}"),
    ("PATCH", "/api/v1/qa-pairs/{qa_pair_id}/status"),
    ("POST", "/api/v1/conversations"),
    ("GET", "/api/v1/conversations"),
    ("GET", "/api/v1/conversations/{conversation_id}/messages"),
    ("DELETE", "/api/v1/conversations/{conversation_id}"),
    ("POST", "/api/v1/retrieval/search"),
    ("GET", "/api/v1/retrieval/logs"),
    ("GET", "/api/v1/retrieval/logs/{log_id}"),
    ("POST", "/api/v1/chat/stream"),
    ("POST", "/api/v1/chat"),
    ("GET", "/api/v1/settings/rag"),
    ("PUT", "/api/v1/settings/rag"),
    ("GET", "/api/v1/evaluation/cases"),
    ("POST", "/api/v1/evaluation/cases"),
    ("POST", "/api/v1/evaluation/runs"),
    ("GET", "/api/v1/evaluation/runs/{run_id}"),
}


def validate_api_coverage() -> None:
    documented = {(item.method, item.path) for item in endpoints()}
    if documented != EXPECTED_OPERATIONS:
        missing = sorted(EXPECTED_OPERATIONS - documented)
        extra = sorted(documented - EXPECTED_OPERATIONS)
        raise RuntimeError(f"API document coverage mismatch. missing={missing}, extra={extra}")

    sys.path.insert(0, str(ROOT / "backend"))
    from app.main import app

    spec = app.openapi()
    actual = set()
    for path, ops in spec["paths"].items():
        for method in ops:
            if method.upper() in {"GET", "POST", "PUT", "PATCH", "DELETE", "HEAD"}:
                actual.add((method.upper(), path))
    if actual != EXPECTED_OPERATIONS:
        missing = sorted(actual - EXPECTED_OPERATIONS)
        extra = sorted(EXPECTED_OPERATIONS - actual)
        raise RuntimeError(f"OpenAPI coverage mismatch. new={missing}, stale={extra}")


def build_api_markdown() -> str:
    validate_api_coverage()
    lines = [
        "# 学校 RAG 智能问答系统接口测试文档",
        "",
        "生成日期：2026-06-12",
        "",
        "## 1. 测试约定",
        "",
        f"- 默认后端地址：`{BASE_URL}`",
        "- 统一接口前缀：`/api/v1`，根路径 `/` 除外。",
        "- 受保护接口需请求头：`Authorization: Bearer $TOKEN`。",
        "- 登录获取 token 示例：",
        "",
        "```bash",
        f"""TOKEN=$(curl -s -X POST {BASE_URL}/api/v1/auth/login \\
  -H "Content-Type: application/json" \\
  -d '{{"username":"admin","password":"admin123"}}' | jq -r '.data.access_token')""",
        "```",
        "",
        "- 统一成功响应：",
        "",
        "```json",
        '{"code":0,"message":"ok","data":{}}',
        "```",
        "",
        "- FastAPI 参数校验错误通常返回 HTTP 422；未登录或 token 无效返回 HTTP 401；业务不存在常见返回 HTTP 404。",
        "- 文件流接口和 SSE 流式接口不使用统一响应包装。",
        "",
        "## 2. 状态枚举与公共对象",
        "",
        "### 2.1 文档状态",
        "",
        md_table(
            ["状态", "含义"],
            [
                ["uploaded", "已上传，尚未完成处理。"],
                ["parsed", "已解析，待切片或向量化。"],
                ["chunked", "已切片，待向量化。"],
                ["indexed", "已完成向量化，可参与检索。"],
                ["failed", "处理失败。"],
                ["needs_conversion", "旧版 Office 文件需先转换。"],
                ["needs_extraction", "压缩包需先解压导入。"],
                ["no_indexable_content", "解析完成但没有可入库正文。"],
                ["converted", "源文件已转换，请查看转换后的子文档。"],
                ["extracted", "压缩包已解压，请查看导入的子文档。"],
                ["deleted", "已软删除。"],
            ],
        ),
        "### 2.2 任务状态",
        "",
        md_table(
            ["状态", "含义"],
            [
                ["pending", "等待执行。"],
                ["running", "执行中。"],
                ["succeeded", "执行成功。"],
                ["failed", "执行失败。"],
                ["skipped", "因状态不满足被跳过。"],
                ["canceled", "已取消，目前代码仅保留枚举。"],
            ],
        ),
        "### 2.3 通用 Document 字段",
        "",
        md_table(["字段", "类型", "必填", "说明"], [[x.name, x.type, x.required, x.desc] for x in DOCUMENT_FIELDS]),
        "### 2.4 通用 Job 字段",
        "",
        md_table(["字段", "类型", "必填", "说明"], [[x.name, x.type, x.required, x.desc] for x in JOB_FIELDS]),
        "## 3. 接口明细",
        "",
    ]
    for index, item in enumerate(endpoints(), start=1):
        lines.extend(
            [
                f"### 3.{index} {item.title}",
                "",
                f"- 方法：`{item.method}`",
                f"- 路径：`{item.path}`",
                f"- 需要登录：{item.auth}",
                f"- 说明：{item.desc}",
                "",
                "请求参数：",
                "",
                md_table(["字段", "类型", "必填", "说明"], [[x.name, x.type, x.required, x.desc] for x in item.params]),
                "请求体：",
                "",
                md_table(["字段", "类型", "必填", "说明"], [[x.name, x.type, x.required, x.desc] for x in item.request]),
                "响应字段：",
                "",
                md_table(["字段", "类型", "必填", "说明"], [[x.name, x.type, x.required, x.desc] for x in item.response]),
                "调用示例：",
                "",
                "```bash",
                item.curl,
                "```",
                "",
            ]
        )
        if item.notes:
            lines.append("注意事项：")
            lines.append("")
            for note in item.notes:
                lines.append(f"- {note}")
            lines.append("")
    return "\n".join(lines).rstrip() + "\n"


def parse_sql_schema() -> tuple[list[dict], dict[str, list[str]], list[str]]:
    sql = (ROOT / "backend" / "sql" / "001_init.sql").read_text(encoding="utf-8")
    extensions = re.findall(r"CREATE EXTENSION IF NOT EXISTS ([^;]+);", sql, flags=re.I)
    tables = []
    for match in re.finditer(r"CREATE TABLE IF NOT EXISTS\s+(\w+)\s*\((.*?)\);", sql, flags=re.S | re.I):
        table_name = match.group(1)
        body = match.group(2)
        columns = []
        for raw in split_sql_columns(body):
            line = " ".join(raw.strip().split())
            if not line or line.upper().startswith(("PRIMARY KEY", "FOREIGN KEY", "CONSTRAINT", "UNIQUE ")):
                continue
            col_match = re.match(r"(\w+)\s+(.+)", line)
            if not col_match:
                continue
            name, rest = col_match.groups()
            columns.append(
                {
                    "name": name,
                    "type": column_type(rest),
                    "nullable": "否" if " NOT NULL" in f" {rest.upper()}" or " PRIMARY KEY" in f" {rest.upper()}" else "是",
                    "default": column_default(rest),
                    "desc": column_desc(table_name, name),
                    "constraint": column_constraint(rest),
                }
            )
        tables.append({"name": table_name, "desc": table_desc(table_name), "columns": columns})
    indexes: dict[str, list[str]] = {table["name"]: [] for table in tables}
    for statement in re.findall(r"CREATE INDEX IF NOT EXISTS\s+.*?;", sql, flags=re.S | re.I):
        compact = " ".join(statement.strip().split())
        table_match = re.search(r"\sON\s+(\w+)\s", compact, flags=re.I)
        if table_match:
            indexes.setdefault(table_match.group(1), []).append(compact)
    return tables, indexes, extensions


def split_sql_columns(body: str) -> list[str]:
    parts = []
    current = []
    depth = 0
    for char in body:
        if char == "(":
            depth += 1
        elif char == ")":
            depth -= 1
        if char == "," and depth == 0:
            parts.append("".join(current))
            current = []
        else:
            current.append(char)
    if current:
        parts.append("".join(current))
    return parts


def column_type(rest: str) -> str:
    tokens = rest.split()
    collected = []
    for token in tokens:
        if token.upper() in {"PRIMARY", "NOT", "NULL", "DEFAULT", "UNIQUE", "REFERENCES", "ON"}:
            break
        collected.append(token)
    return " ".join(collected)


def column_default(rest: str) -> str:
    match = re.search(r"\bDEFAULT\s+(.+?)(?:\s+REFERENCES|\s+NOT\s+NULL|\s+NULL|\s+PRIMARY\s+KEY|\s+UNIQUE|$)", rest, flags=re.I)
    return match.group(1).strip() if match else ""


def column_constraint(rest: str) -> str:
    items = []
    upper = rest.upper()
    if "PRIMARY KEY" in upper:
        items.append("主键")
    if "UNIQUE" in upper:
        items.append("唯一")
    ref = re.search(r"REFERENCES\s+([^)]+(?:\)[^ ]*)?)", rest, flags=re.I)
    if ref:
        items.append(f"外键 {ref.group(1).strip()}")
    return "；".join(items)


TABLE_DESCS = {
    "admins": "管理员账号表。",
    "documents": "文档主表，保存文件元数据、存储位置和处理状态。",
    "document_jobs": "文档后台任务表，记录解析、切片、向量化、转换、解压、QA 生成等任务。",
    "document_parse_results": "文档解析结果表，一篇文档最多一条解析结果。",
    "document_chunks": "文档切片表，保存可检索文本片段和父子切片关系。",
    "chunk_embeddings": "文档切片向量表，使用 pgvector 存储 1024 维 embedding。",
    "qa_pairs": "人工或自动生成的问答对表。",
    "qa_pair_embeddings": "QA 对向量表，使用 pgvector 存储 1024 维 embedding。",
    "conversations": "聊天会话表。",
    "conversation_messages": "聊天消息表，保存问答内容、引用和检索追踪。",
    "retrieval_logs": "检索与回答日志表，用于调试和评估 RAG 效果。",
    "evaluation_cases": "评测用例表。",
    "evaluation_runs": "评测批次表。",
    "evaluation_results": "评测结果明细表。",
}


COLUMN_DESCS = {
    "id": "主键 UUID。",
    "created_at": "创建时间。",
    "updated_at": "更新时间。",
    "deleted_at": "软删除时间，非空表示已删除。",
    "username": "管理员登录账号。",
    "password_hash": "管理员密码哈希。",
    "display_name": "管理员显示名称。",
    "is_active": "是否启用。",
    "last_login_at": "最后登录时间。",
    "title": "标题。",
    "file_name": "原始文件名。",
    "file_ext": "文件扩展名。",
    "mime_type": "MIME 类型。",
    "file_size": "文件大小，单位字节。",
    "file_hash": "文件 SHA256 哈希。",
    "storage_bucket": "MinIO bucket。",
    "storage_object_key": "MinIO object key。",
    "source_url": "外部来源 URL。",
    "preview_url": "预览 URL。",
    "download_url": "下载 URL。",
    "status": "业务状态。",
    "parse_quality_score": "解析质量分。",
    "error_message": "错误信息。",
    "created_by": "创建人管理员 ID。",
    "updated_by": "更新人管理员 ID。",
    "document_id": "关联文档 ID。",
    "job_type": "任务类型。",
    "progress": "任务进度百分比。",
    "message": "任务消息。",
    "retry_count": "重试次数。",
    "params": "任务参数 JSON。",
    "result": "任务结果 JSON。",
    "started_at": "任务开始时间。",
    "finished_at": "任务结束时间。",
    "parser_name": "解析器名称。",
    "parser_version": "解析器版本。",
    "content_text": "解析纯文本。",
    "content_md": "解析 Markdown。",
    "content_object_key": "解析内容在 MinIO 中的 object key。",
    "page_count": "页数。",
    "quality_score": "质量分。",
    "parse_meta": "解析元数据 JSON。",
    "parent_chunk_id": "父切片 ID。",
    "chunk_no": "切片序号。",
    "chunk_type": "切片类型。",
    "content": "文本内容。",
    "content_hash": "内容哈希。",
    "token_count": "token 数。",
    "char_count": "字符数。",
    "page_start": "起始页码。",
    "page_end": "结束页码。",
    "section_path": "章节路径。",
    "metadata": "元数据 JSON。",
    "search_vector": "PostgreSQL 全文检索向量。",
    "embedding": "pgvector 向量。",
    "embedding_model": "向量模型名称。",
    "embedding_provider": "向量提供方。",
    "question": "问题。",
    "answer": "答案。",
    "source_document_id": "来源文档 ID。",
    "source_chunk_ids": "来源切片 ID 数组。",
    "tags": "标签数组。",
    "version": "版本号。",
    "qa_pair_id": "QA 对 ID。",
    "summary": "会话摘要。",
    "message_count": "消息数。",
    "last_message_at": "最后消息时间。",
    "conversation_id": "会话 ID。",
    "role": "消息角色。",
    "rewritten_query": "改写后的检索问题。",
    "retrieval_trace": "检索追踪 JSON。",
    "citations": "引用来源 JSON。",
    "suggested_questions": "推荐追问 JSON。",
    "latency_ms": "耗时毫秒。",
    "model_name": "模型名称。",
    "message_id": "消息 ID。",
    "raw_query": "原始查询。",
    "recall_results": "召回结果 JSON。",
    "rerank_results": "重排结果 JSON。",
    "final_context": "最终上下文 JSON。",
    "answer_quality": "回答质量 JSON。",
    "rerank_model": "重排模型名称。",
    "expected_answer": "期望答案或关键词。",
    "expected_document_ids": "期望命中文档 ID 数组。",
    "expected_chunk_ids": "期望命中切片 ID 数组。",
    "difficulty": "难度。",
    "name": "名称。",
    "config": "配置 JSON。",
    "metrics": "指标 JSON。",
    "run_id": "评测批次 ID。",
    "case_id": "评测用例 ID。",
    "hit_top3": "Top3 是否命中。",
    "hit_top5": "Top5 是否命中。",
    "answer_score": "答案得分。",
    "citation_score": "引用得分。",
    "suggested_question_score": "推荐追问得分。",
    "failure_reason": "失败原因。",
    "trace": "追踪 JSON。",
}


def table_desc(table: str) -> str:
    return TABLE_DESCS.get(table, "")


def column_desc(table: str, column: str) -> str:
    if table == "evaluation_runs" and column == "status":
        return "评测批次状态：running、succeeded、failed。"
    if table == "document_jobs" and column == "status":
        return "任务状态：pending、running、succeeded、failed、skipped、canceled。"
    if table == "documents" and column == "status":
        return "文档状态，见接口文档状态枚举。"
    if table == "qa_pairs" and column == "status":
        return "QA 状态，常用 enabled、disabled、draft。"
    return COLUMN_DESCS.get(column, "")


def build_db_markdown() -> str:
    tables, indexes, extensions = parse_sql_schema()
    lines = [
        "# 学校 RAG 智能问答系统数据库表结构文档",
        "",
        "生成日期：2026-06-12",
        "",
        "## 1. 数据库概览",
        "",
        "- 数据库：PostgreSQL。",
        "- 主要扩展：pgvector、pg_trgm、uuid-ossp。",
        "- 向量维度：1024，表 `chunk_embeddings` 和 `qa_pair_embeddings` 使用 HNSW 索引。",
        "- 本文档依据当前 `backend/sql/001_init.sql` 生成。",
        "",
        "扩展清单：",
        "",
        md_table(["扩展"], [[x.strip()] for x in extensions]),
        "## 2. 表关系概要",
        "",
        md_table(
            ["表名", "说明"],
            [[table["name"], table["desc"]] for table in tables],
        ),
        "## 3. 表字段明细",
        "",
    ]
    for index, table in enumerate(tables, start=1):
        lines.extend(
            [
                f"### 3.{index} {table['name']}",
                "",
                table["desc"],
                "",
                md_table(
                    ["字段", "类型", "允许为空", "默认值", "约束", "说明"],
                    [
                        [
                            col["name"],
                            col["type"],
                            col["nullable"],
                            col["default"],
                            col["constraint"],
                            col["desc"],
                        ]
                        for col in table["columns"]
                    ],
                ),
                "索引：",
                "",
                md_table(["索引定义"], [[idx] for idx in indexes.get(table["name"], [])]),
            ]
        )
    lines.extend(
        [
            "## 4. 关键业务关系",
            "",
            "- `documents` 是文档主表，`document_parse_results`、`document_chunks`、`chunk_embeddings` 均通过 `document_id` 关联文档。",
            "- `document_parse_results.document_id` 唯一，一篇文档最多一条解析结果。",
            "- `document_chunks.parent_chunk_id` 指向同表，用于父子切片关系。",
            "- `chunk_embeddings.chunk_id` 唯一，一个切片最多一条 embedding。",
            "- `qa_pairs` 可选关联 `documents`，`qa_pair_embeddings.qa_pair_id` 唯一。",
            "- `conversations` 与 `conversation_messages` 是一对多关系。",
            "- `retrieval_logs` 可选关联会话和消息，用于审计和调试。",
            "- `evaluation_runs`、`evaluation_cases`、`evaluation_results` 构成评测批次、用例和结果明细关系。",
            "",
        ]
    )
    return "\n".join(lines).rstrip() + "\n"


def add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    if not rows:
        doc.add_paragraph("无")
        return
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for idx, header in enumerate(headers):
        table.rows[0].cells[idx].text = header
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].text = value


def add_code(doc: Document, code: str) -> None:
    paragraph = doc.add_paragraph()
    run = paragraph.add_run(code)
    run.font.name = "Consolas"
    run.font.size = Pt(9)


def build_api_docx() -> None:
    doc = Document()
    doc.add_heading("学校 RAG 智能问答系统接口测试文档", 0)
    doc.add_paragraph("生成日期：2026-06-12")
    doc.add_heading("1. 测试约定", 1)
    for text in [
        f"默认后端地址：{BASE_URL}",
        "统一接口前缀：/api/v1，根路径 / 除外。",
        "受保护接口需请求头：Authorization: Bearer $TOKEN。",
        "统一成功响应：{\"code\":0,\"message\":\"ok\",\"data\":{}}。",
        "文件流接口和 SSE 流式接口不使用统一响应包装。",
    ]:
        doc.add_paragraph(text, style="List Bullet")
    doc.add_paragraph("登录获取 token 示例：")
    add_code(doc, f"""TOKEN=$(curl -s -X POST {BASE_URL}/api/v1/auth/login \\
  -H "Content-Type: application/json" \\
  -d '{{"username":"admin","password":"admin123"}}' | jq -r '.data.access_token')""")
    doc.add_heading("2. 状态枚举与公共对象", 1)
    doc.add_heading("2.1 文档状态", 2)
    add_table(
        doc,
        ["状态", "含义"],
        [
            ["uploaded", "已上传，尚未完成处理。"],
            ["parsed", "已解析，待切片或向量化。"],
            ["chunked", "已切片，待向量化。"],
            ["indexed", "已完成向量化，可参与检索。"],
            ["failed", "处理失败。"],
            ["needs_conversion", "旧版 Office 文件需先转换。"],
            ["needs_extraction", "压缩包需先解压导入。"],
            ["no_indexable_content", "解析完成但没有可入库正文。"],
            ["converted", "源文件已转换，请查看转换后的子文档。"],
            ["extracted", "压缩包已解压，请查看导入的子文档。"],
            ["deleted", "已软删除。"],
        ],
    )
    doc.add_heading("2.2 通用 Document 字段", 2)
    add_table(doc, ["字段", "类型", "必填", "说明"], [[x.name, x.type, x.required, x.desc] for x in DOCUMENT_FIELDS])
    doc.add_heading("2.3 通用 Job 字段", 2)
    add_table(doc, ["字段", "类型", "必填", "说明"], [[x.name, x.type, x.required, x.desc] for x in JOB_FIELDS])
    doc.add_heading("3. 接口明细", 1)
    for index, item in enumerate(endpoints(), start=1):
        doc.add_heading(f"3.{index} {item.title}", 2)
        for text in [
            f"方法：{item.method}",
            f"路径：{item.path}",
            f"需要登录：{item.auth}",
            f"说明：{item.desc}",
        ]:
            doc.add_paragraph(text)
        doc.add_paragraph("请求参数：")
        add_table(doc, ["字段", "类型", "必填", "说明"], [[x.name, x.type, x.required, x.desc] for x in item.params])
        doc.add_paragraph("请求体：")
        add_table(doc, ["字段", "类型", "必填", "说明"], [[x.name, x.type, x.required, x.desc] for x in item.request])
        doc.add_paragraph("响应字段：")
        add_table(doc, ["字段", "类型", "必填", "说明"], [[x.name, x.type, x.required, x.desc] for x in item.response])
        doc.add_paragraph("调用示例：")
        add_code(doc, item.curl)
        if item.notes:
            doc.add_paragraph("注意事项：")
            for note in item.notes:
                doc.add_paragraph(note, style="List Bullet")
    doc.save(API_DOCX)


def build_db_docx() -> None:
    tables, indexes, extensions = parse_sql_schema()
    doc = Document()
    doc.add_heading("学校 RAG 智能问答系统数据库表结构文档", 0)
    doc.add_paragraph("生成日期：2026-06-12")
    doc.add_heading("1. 数据库概览", 1)
    for text in [
        "数据库：PostgreSQL。",
        "主要扩展：pgvector、pg_trgm、uuid-ossp。",
        "向量维度：1024，表 chunk_embeddings 和 qa_pair_embeddings 使用 HNSW 索引。",
        "本文档依据当前 backend/sql/001_init.sql 生成。",
    ]:
        doc.add_paragraph(text, style="List Bullet")
    doc.add_heading("扩展清单", 2)
    add_table(doc, ["扩展"], [[x.strip()] for x in extensions])
    doc.add_heading("2. 表关系概要", 1)
    add_table(doc, ["表名", "说明"], [[table["name"], table["desc"]] for table in tables])
    doc.add_heading("3. 表字段明细", 1)
    for index, table in enumerate(tables, start=1):
        doc.add_heading(f"3.{index} {table['name']}", 2)
        doc.add_paragraph(table["desc"])
        add_table(
            doc,
            ["字段", "类型", "允许为空", "默认值", "约束", "说明"],
            [
                [col["name"], col["type"], col["nullable"], col["default"], col["constraint"], col["desc"]]
                for col in table["columns"]
            ],
        )
        doc.add_paragraph("索引：")
        add_table(doc, ["索引定义"], [[idx] for idx in indexes.get(table["name"], [])])
    doc.add_heading("4. 关键业务关系", 1)
    for text in [
        "documents 是文档主表，document_parse_results、document_chunks、chunk_embeddings 均通过 document_id 关联文档。",
        "document_parse_results.document_id 唯一，一篇文档最多一条解析结果。",
        "document_chunks.parent_chunk_id 指向同表，用于父子切片关系。",
        "chunk_embeddings.chunk_id 唯一，一个切片最多一条 embedding。",
        "qa_pairs 可选关联 documents，qa_pair_embeddings.qa_pair_id 唯一。",
        "conversations 与 conversation_messages 是一对多关系。",
        "retrieval_logs 可选关联会话和消息，用于审计和调试。",
        "evaluation_runs、evaluation_cases、evaluation_results 构成评测批次、用例和结果明细关系。",
    ]:
        doc.add_paragraph(text, style="List Bullet")
    doc.save(DB_DOCX)


def main() -> None:
    DOCS_DIR.mkdir(parents=True, exist_ok=True)
    API_MD.write_text(build_api_markdown(), encoding="utf-8")
    DB_MD.write_text(build_db_markdown(), encoding="utf-8")
    build_api_docx()
    build_db_docx()
    print(API_MD)
    print(API_DOCX)
    print(DB_MD)
    print(DB_DOCX)


if __name__ == "__main__":
    main()
