import json
import re
import shutil
import subprocess
import zipfile
import xml.etree.ElementTree as ET
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from typing import Any

from app.core.config import get_settings
from app.services.ocr_service import OcrResult, PaddleOcrService


class ParsedDocument:
    def __init__(
        self,
        text: str,
        markdown: str,
        page_count: int | None,
        meta: dict,
        assets: list[dict] | None = None,
    ) -> None:
        self.text = text
        self.markdown = markdown
        self.page_count = page_count
        self.meta = meta
        self.assets = assets or []


class ParseService:
    IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".bmp", ".tif", ".tiff"}
    CONTAINER_EXTENSIONS = {".zip", ".rar"}

    def __init__(self) -> None:
        self.settings = get_settings()
        self.ocr_service = PaddleOcrService(
            lang=self.settings.ocr_lang,
            use_gpu=self.settings.ocr_use_gpu,
            min_confidence=self.settings.ocr_min_confidence,
        )

    def parse(self, file_name: str, data: bytes) -> ParsedDocument:
        ext = Path(file_name).suffix.lower()
        if ext == ".pdf":
            result = self._parse_pdf(data)
        elif ext in {".txt", ".md"}:
            text = data.decode("utf-8", errors="ignore")
            result = ParsedDocument(text=text, markdown=text, page_count=None, meta={"parser": "plain-text"})
        elif ext == ".docx":
            result = self._parse_docx(data)
        elif ext == ".doc":
            result = self._parse_doc(data)
        elif ext in {".xlsx", ".xlsm"}:
            result = self._parse_xlsx(data)
        elif ext == ".xls":
            result = self._parse_xls_with_fallback(data)
        elif ext in self.IMAGE_EXTENSIONS:
            result = self._parse_image(file_name, data)
        elif ext in self.CONTAINER_EXTENSIONS:
            result = self._parse_container(file_name, data)
        else:
            raise ValueError(f"暂不支持的文件类型: {ext}")
        # 统一应用清洗
        cleaned_text = self._clean_parsed_text(result.text)
        cleaned_md = self._clean_parsed_text(result.markdown or "")
        return ParsedDocument(
            text=cleaned_text,
            markdown=cleaned_md or cleaned_text,
            page_count=result.page_count,
            meta=result.meta,
            assets=result.assets,
        )

    @staticmethod
    def _clean_parsed_text(text: str) -> str:
        lines = text.splitlines()
        cleaned: list[str] = []
        page_footer_patterns = [
            r"^\d+\s*/\s*\d+$",           # "1 / 10"
            r"^\d+\s*of\s*\d+$",           # "1 of 10"
            r"^第\s*\d+\s*页$",            # "第 1 页"
            r"^-\s*\d+\s*-$",              # "- 1 -"
            r"^--+\s*第\s*\d+\s*页\s*--+$", # "----第1页----"
            r"^\[\d+\]$",                  # "[1]"
        ]
        header_footer_pattern = re.compile("|".join(page_footer_patterns))

        # 常见的页眉关键词（出现在文档每页顶部）
        header_keywords = [
            "庄艳蓓",   # 合同中的场景
        ]

        # 连续相同行（目录重复）检测
        prev_line = None
        prev_repeat_count = 0
        blank_streak = 0

        for raw_line in lines:
            line = raw_line.strip()

            # 空行压缩（连续空行最多保留1个）
            if not line:
                if blank_streak > 0:
                    continue
                blank_streak += 1
                cleaned.append("")
                prev_line = None
                continue
            blank_streak = 0

            # 纯数字/页码行
            if header_footer_pattern.match(line):
                continue

            # 单字符无意义行
            if len(line) <= 1:
                continue

            # 纯符号行（---、===、*** 等）
            if re.fullmatch(r"[=\-*_~#\s]{2,}", line):
                continue

            # 页眉关键词 + 短文本（"庄艳蓓 2026-01-01" 之类的页眉行）
            if any(kw in line for kw in header_keywords) and len(line) < 40:
                continue

            # 检测目录重复行（连续三行以上相同的内容）
            if prev_line and line == prev_line:
                prev_repeat_count += 1
                if prev_repeat_count >= 2:
                    # 跳过这条重复行（通常是目录重复）
                    continue
            else:
                prev_repeat_count = 0
            prev_line = line

            cleaned.append(line)

        # 尾部多余空行
        while cleaned and cleaned[-1] == "":
            cleaned.pop()

        result = "\n".join(cleaned)

        # 移除连续"目录"章节中"..."和页码的冗余行
        result = re.sub(r"\.{2,}\s*\d+\s*$", "", result, flags=re.MULTILINE)
        # 移除类似 "............. 1" 的行
        result = re.sub(r"^[\.\s]{4,}\d{1,3}\s*$", "", result, flags=re.MULTILINE)

        return result.strip()

    def _parse_pdf(self, data: bytes) -> ParsedDocument:
        import fitz

        doc = fitz.open(stream=data, filetype="pdf")
        pages: list[str] = []
        md_pages: list[str] = []
        scanned_pages: list[int] = []
        low_text_pages: list[int] = []
        ocr_pages: list[dict] = []
        ocr_errors: list[dict] = []
        for index, page in enumerate(doc, start=1):
            text = page.get_text("text").strip()
            needs_ocr = self._pdf_text_needs_ocr(page, text)
            if needs_ocr:
                if text:
                    low_text_pages.append(index)
                else:
                    scanned_pages.append(index)

                ocr_result = (
                    self._ocr_pdf_page(page)
                    if self._should_ocr(len(ocr_pages) + len(ocr_errors))
                    else None
                )
                if ocr_result and ocr_result.text and self._prefer_ocr_text(text, ocr_result.text):
                    page_text = ocr_result.text
                    block = self._rag_block(
                        "ocr_page",
                        {
                            "page": index,
                            "source": "pdf",
                            **self._ocr_summary(ocr_result),
                        },
                        page_text,
                    )
                    pages.append(f"[第{index}页]\n{block}")
                    md_pages.append(f"\n\n## 第{index}页\n\n{block}")
                    ocr_pages.append({"page": index, **self._ocr_summary(ocr_result)})
                    continue

                if ocr_result and ocr_result.error:
                    ocr_errors.append({"page": index, "error": ocr_result.error})
                elif ocr_result is None:
                    ocr_errors.append({"page": index, "error": self._ocr_skip_reason()})

                if not text:
                    error = ocr_errors[-1]["error"] if ocr_errors else "OCR 未返回可用文字"
                    placeholder = f"[本页未提取到文本，PaddleOCR 未返回可用文字：{error}]"
                    pages.append(f"[第{index}页]\n{placeholder}")
                    md_pages.append(f"\n\n## 第{index}页\n\n> {placeholder}")
                    continue

            pages.append(f"[第{index}页]\n{text}")
            md_pages.append(f"\n\n## 第{index}页\n\n{text}")

        full_text = "\n\n".join(pages)
        markdown = self._strip_rag_markers("\n".join(md_pages).strip())
        ocr_needed_count = len(scanned_pages) + len(low_text_pages)
        return ParsedDocument(
            text=full_text,
            markdown=markdown,
            page_count=doc.page_count,
            meta={
                "parser": "pymupdf",
                "page_count": doc.page_count,
                "scanned_page_count": len(scanned_pages),
                "scanned_pages": scanned_pages,
                "low_text_page_count": len(low_text_pages),
                "low_text_pages": low_text_pages,
                "ocr_enabled": self.settings.ocr_enabled,
                "ocr_engine": "paddleocr",
                "ocr_page_count": len(ocr_pages),
                "ocr_pages": ocr_pages,
                "ocr_errors": ocr_errors,
                "ocr_required": len(ocr_pages) < ocr_needed_count,
            },
        )

    def _parse_docx(self, data: bytes) -> ParsedDocument:
        from docx import Document as DocxDocument
        from docx.oxml.table import CT_Tbl
        from docx.oxml.text.paragraph import CT_P
        from docx.table import Table
        from docx.text.paragraph import Paragraph

        try:
            doc = DocxDocument(BytesIO(data))
        except zipfile.BadZipFile as exc:
            return self._parse_docx_raw_xml(data, str(exc))

        parts: list[str] = []
        table_meta: list[dict] = []
        table_index = 0
        paragraph_count = 0

        for child in doc.element.body.iterchildren():
            if isinstance(child, CT_P):
                paragraph = Paragraph(child, doc)
                text = paragraph.text.strip()
                if not text:
                    continue
                paragraph_count += 1
                level = self._docx_heading_level(paragraph)
                parts.append(f"{'#' * min(level + 1, 6)} {text}" if level else text)
            elif isinstance(child, CT_Tbl):
                table_index += 1
                table = Table(child, doc)
                rows = self._docx_table_rows(table)
                if not rows:
                    continue
                metadata = {
                    "source": "docx",
                    "table_index": table_index,
                    "row_start": 1,
                    "row_end": len(rows),
                    "header_row": 1,
                    "row_count": len(rows),
                    "column_count": max((len(row) for row in rows), default=0),
                    "fields": [value for value in rows[0] if value] if rows else [],
                    "index_mode": "full",
                    "indexed_row_count": len(rows),
                    "full_index_limited": False,
                }
                table_rows = [(index, row) for index, row in enumerate(rows, start=1)]
                indexed_rows = table_rows
                if self._is_large_table(table_rows):
                    sample_limit = max(1, int(self.settings.table_sample_rows))
                    indexed_rows = table_rows[:sample_limit]
                    metadata = {
                        **metadata,
                        "index_mode": "sampled",
                        "indexed_row_count": len(indexed_rows),
                        "full_index_limited": True,
                    }
                    parts.extend(
                        self._large_table_summary_lines(
                            f"表格 {table_index}",
                            table_rows,
                            indexed_rows,
                            metadata,
                        )
                    )
                table_meta.append(metadata)
                table_text = f"### 表格 {table_index}\n{self._markdown_table([values for _, values in indexed_rows])}"
                parts.append(self._rag_block("table", metadata, table_text))

        image_blocks: list[str] = []
        image_ocr_results: list[dict] = []
        image_ocr_errors: list[dict] = []
        image_assets: list[dict] = []
        image_index = 0
        image_count = sum(
            1
            for part in doc.part.related_parts.values()
            if getattr(part, "content_type", "").startswith("image/")
        )
        for part in doc.part.related_parts.values():
            content_type = getattr(part, "content_type", "")
            if not content_type.startswith("image/"):
                continue
            image_index += 1
            try:
                image_bytes = part.blob
            except Exception as exc:
                image_ocr_errors.append({"image": image_index, "error": str(exc)})
                continue
            block, result_meta, error_meta, asset = self._build_image_block(
                source="docx",
                image_index=image_index,
                image_bytes=image_bytes,
                content_type=content_type,
                file_name=Path(str(getattr(part, "partname", ""))).name or f"image{image_index}",
                context_text="",
                position_meta={},
                attempted_count=len(image_ocr_results) + len(image_ocr_errors),
            )
            image_blocks.append(block)
            image_assets.append(asset)
            if result_meta:
                image_ocr_results.append(result_meta)
            if error_meta:
                image_ocr_errors.append(error_meta)

        parts.extend(image_blocks)
        if image_count and not image_blocks:
            parts.append(f"[文档包含 {image_count} 张图片，PaddleOCR 未返回可用文字。]")

        text = "\n\n".join(parts)
        return ParsedDocument(
            text=text,
            markdown=self._strip_rag_markers(text),
            page_count=None,
            meta={
                "parser": "python-docx",
                "paragraph_count": paragraph_count,
                "table_count": len(table_meta),
                "tables": table_meta,
                "image_count": image_count,
                "ocr_enabled": self.settings.ocr_enabled,
                "ocr_engine": "paddleocr",
                "image_ocr_count": len(image_ocr_results),
                "image_ocr_results": image_ocr_results,
                "image_ocr_errors": image_ocr_errors,
                "image_ocr_required": len(image_ocr_results) < image_count,
            },
            assets=image_assets,
        )

    def _parse_docx_raw_xml(self, data: bytes, fallback_reason: str) -> ParsedDocument:
        ns = {
            "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
            "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
            "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
        }
        parts: list[str] = []
        table_meta: list[dict] = []
        image_blocks: list[str] = []
        image_ocr_results: list[dict] = []
        image_ocr_errors: list[dict] = []
        image_assets: list[dict] = []
        paragraph_count = 0
        table_index = 0
        image_index = 0
        bad_media: list[dict] = []
        recent_texts: list[str] = []

        with zipfile.ZipFile(BytesIO(data)) as archive:
            rels = self._docx_relationships(archive)
            root = ET.fromstring(archive.read("word/document.xml"))
            body = root.find("w:body", ns)
            if body is None:
                raise RuntimeError("DOCX 缺少 word/document.xml 正文内容")

            for child in list(body):
                if child.tag == f"{{{ns['w']}}}p":
                    text = self._docx_xml_text(child, ns).strip()
                    if text:
                        paragraph_count += 1
                        parts.append(text)
                        recent_texts.append(text)
                        recent_texts = recent_texts[-3:]
                    image_refs = self._docx_xml_image_refs(child, ns, rels)
                    for image_ref in image_refs:
                        image_index += 1
                        target = image_ref.get("target") or ""
                        member_name = self._docx_member_name(target)
                        context_text = " ".join(recent_texts)[-500:]
                        try:
                            image_bytes = self._read_zip_member_crc_tolerant(archive, member_name)
                        except Exception as exc:
                            bad_media.append(
                                {
                                    "image": image_index,
                                    "target": target,
                                    "error": str(exc),
                                }
                            )
                            image_ocr_errors.append({"image": image_index, "error": str(exc)})
                            continue
                        block, result_meta, error_meta, asset = self._build_image_block(
                            source="docx-raw-xml",
                            image_index=image_index,
                            image_bytes=image_bytes,
                            content_type=self._image_content_type_from_name(member_name),
                            file_name=Path(member_name).name,
                            context_text=context_text,
                            position_meta={
                                "paragraph_index": paragraph_count,
                                "relationship_id": image_ref.get("relationship_id"),
                            },
                            attempted_count=len(image_ocr_results) + len(image_ocr_errors),
                        )
                        image_blocks.append(block)
                        parts.append(block)
                        image_assets.append(asset)
                        if result_meta:
                            image_ocr_results.append(result_meta)
                        if error_meta:
                            image_ocr_errors.append(error_meta)
                elif child.tag == f"{{{ns['w']}}}tbl":
                    table_index += 1
                    rows = self._docx_xml_table_rows(child, ns)
                    if not rows:
                        continue
                    metadata = {
                        "source": "docx-raw-xml",
                        "table_index": table_index,
                        "row_start": 1,
                        "row_end": len(rows),
                        "header_row": 1,
                        "row_count": len(rows),
                        "column_count": max((len(row) for row in rows), default=0),
                        "fields": [value for value in rows[0] if value] if rows else [],
                        "index_mode": "full",
                        "indexed_row_count": len(rows),
                        "full_index_limited": False,
                    }
                    table_meta.append(metadata)
                    table_text = f"### 表格 {table_index}\n{self._markdown_table(rows)}"
                    parts.append(self._rag_block("table", metadata, table_text))

        if image_index and not image_blocks:
            parts.append(f"[文档包含 {image_index} 张图片，但图片数据不可读取或 OCR 未返回可用文字。]")

        text = "\n\n".join(parts)
        return ParsedDocument(
            text=text,
            markdown=self._strip_rag_markers(text),
            page_count=None,
            meta={
                "parser": "docx-raw-xml",
                "fallback_from": "python-docx",
                "fallback_reason": fallback_reason,
                "paragraph_count": paragraph_count,
                "table_count": len(table_meta),
                "tables": table_meta,
                "image_count": image_index,
                "bad_media": bad_media,
                "ocr_enabled": self.settings.ocr_enabled,
                "ocr_engine": "paddleocr",
                "image_ocr_count": len(image_ocr_results),
                "image_ocr_results": image_ocr_results,
                "image_ocr_errors": image_ocr_errors,
                "image_ocr_required": len(image_ocr_results) < image_index,
            },
            assets=image_assets,
        )

    def _docx_relationships(self, archive: zipfile.ZipFile) -> dict[str, str]:
        try:
            raw = archive.read("word/_rels/document.xml.rels")
        except KeyError:
            return {}
        root = ET.fromstring(raw)
        rels: dict[str, str] = {}
        for item in root:
            rel_id = item.attrib.get("Id")
            target = item.attrib.get("Target")
            if rel_id and target:
                rels[rel_id] = target
        return rels

    def _docx_xml_text(self, element: ET.Element, ns: dict[str, str]) -> str:
        parts: list[str] = []
        for node in element.iter():
            if node.tag == f"{{{ns['w']}}}t":
                parts.append(node.text or "")
            elif node.tag == f"{{{ns['w']}}}tab":
                parts.append("\t")
            elif node.tag == f"{{{ns['w']}}}br":
                parts.append("\n")
        return "".join(parts)

    def _docx_xml_table_rows(self, table: ET.Element, ns: dict[str, str]) -> list[list[str]]:
        rows: list[list[str]] = []
        for row in table.findall("w:tr", ns):
            values = []
            for cell in row.findall("w:tc", ns):
                text = re.sub(r"\s+", " ", self._docx_xml_text(cell, ns)).strip()
                values.append(text)
            while values and not values[-1]:
                values.pop()
            if any(values):
                rows.append(values)
        return rows

    def _docx_xml_image_refs(
        self,
        paragraph: ET.Element,
        ns: dict[str, str],
        rels: dict[str, str],
    ) -> list[dict]:
        refs: list[dict] = []
        embed_key = f"{{{ns['r']}}}embed"
        for blip in paragraph.findall(".//a:blip", ns):
            rel_id = blip.attrib.get(embed_key)
            refs.append(
                {
                    "relationship_id": rel_id,
                    "target": rels.get(rel_id or "", ""),
                }
            )
        return refs

    def _docx_member_name(self, target: str) -> str:
        normalized = target.replace("\\", "/")
        if normalized.startswith("/"):
            return normalized.lstrip("/")
        if normalized.startswith("word/"):
            return normalized
        return f"word/{normalized.lstrip('./')}"

    def _read_zip_member_crc_tolerant(self, archive: zipfile.ZipFile, member_name: str) -> bytes:
        try:
            return archive.read(member_name)
        except zipfile.BadZipFile:
            with archive.open(member_name) as stream:
                if hasattr(stream, "_expected_crc"):
                    stream._expected_crc = None
                return stream.read()

    def _image_content_type_from_name(self, name: str) -> str:
        suffix = Path(name).suffix.lower()
        if suffix in {".jpg", ".jpeg"}:
            return "image/jpeg"
        if suffix == ".png":
            return "image/png"
        if suffix == ".bmp":
            return "image/bmp"
        if suffix in {".tif", ".tiff"}:
            return "image/tiff"
        return "application/octet-stream"

    def _build_image_block(
        self,
        source: str,
        image_index: int,
        image_bytes: bytes,
        content_type: str,
        file_name: str,
        context_text: str,
        position_meta: dict,
        attempted_count: int,
    ) -> tuple[str, dict | None, dict | None, dict]:
        image_id = f"img_{image_index}"
        suffix = Path(file_name).suffix.lower() or self._image_suffix(content_type)
        safe_file_name = f"{image_id}{suffix}"
        asset = {
            "image_id": image_id,
            "image_index": image_index,
            "file_name": safe_file_name,
            "content_type": content_type,
            "context_text": context_text,
            "data": image_bytes,
            **position_meta,
        }
        metadata = {
            "source": source,
            "image_id": image_id,
            "image_index": image_index,
            "image_file_name": safe_file_name,
            "image_content_type": content_type,
            "image_context": context_text,
            "image_bucket": f"__RAG_IMAGE_BUCKET_{image_id}__",
            "image_object_key": f"__RAG_IMAGE_OBJECT_KEY_{image_id}__",
            **position_meta,
        }
        ocr_result: OcrResult | None = None
        error_meta: dict | None = None
        if self._should_ocr(attempted_count):
            try:
                ocr_result = self.ocr_service.ocr_image_bytes(image_bytes, suffix)
            except Exception as exc:
                error_meta = {"image": image_index, "image_id": image_id, "error": str(exc)}
        else:
            error_meta = {"image": image_index, "image_id": image_id, "error": self._ocr_skip_reason()}

        lines = [f"### 图片 {image_index} OCR"]
        if context_text:
            lines.append(f"相邻文本：{context_text}")
        result_meta: dict | None = None
        if ocr_result and ocr_result.text:
            summary = self._ocr_summary(ocr_result)
            metadata.update(summary)
            lines.append(ocr_result.text)
            result_meta = {"image": image_index, "image_id": image_id, **summary}
        else:
            if ocr_result and not error_meta:
                error_meta = {
                    "image": image_index,
                    "image_id": image_id,
                    "error": ocr_result.error or "PaddleOCR 未返回可用文字",
                }
            lines.append("[图片暂无可用 OCR 文字]")

        return self._rag_block("image_ocr", metadata, "\n".join(lines)), result_meta, error_meta, asset

    def _parse_doc(self, data: bytes) -> ParsedDocument:
        converter = self._office_converter()
        converted = self._convert_legacy_office(data, ".doc", "docx")
        if converted:
            parsed = self._parse_docx(converted)
            parsed.meta = {
                **parsed.meta,
                "parser": "libreoffice+python-docx",
                "converted_from": ".doc",
            }
            return parsed
        if converter:
            raise RuntimeError("LibreOffice 转换旧版 Word .doc 失败，文件可能损坏、加密或格式不兼容")
        return self._needs_conversion_document(
            ".doc",
            "旧版 Word .doc 需要先通过 LibreOffice 或 Microsoft Word 转换为 .docx。",
        )

    def _parse_xlsx(self, data: bytes) -> ParsedDocument:
        from io import BytesIO

        from openpyxl import load_workbook

        workbook = load_workbook(BytesIO(data), data_only=True)
        sections: list[str] = []
        sheet_meta: list[dict] = []
        image_blocks: list[str] = []
        image_ocr_results: list[dict] = []
        image_ocr_errors: list[dict] = []
        image_count = 0
        image_index = 0
        for sheet_index, sheet in enumerate(workbook.worksheets, start=1):
            rows = self._openpyxl_rows(sheet)
            merged_ranges = [str(item) for item in sheet.merged_cells.ranges]
            sheet_text, metadata = self._excel_sheet_markdown(
                sheet.title,
                rows,
                merged_ranges,
                parser="openpyxl",
                sheet_index=sheet_index,
            )
            sections.append(sheet_text)
            sheet_meta.append(metadata)
            blocks, results, errors, count, image_index = self._openpyxl_image_ocr_blocks(
                sheet,
                sheet_index,
                image_index,
            )
            image_blocks.extend(blocks)
            image_ocr_results.extend(results)
            image_ocr_errors.extend(errors)
            image_count += count
        sections.extend(image_blocks)
        if image_count and not image_blocks:
            sections.append(f"[工作簿包含 {image_count} 张图片，PaddleOCR 未返回可用文字。]")
        text = "\n\n".join(sections)
        return ParsedDocument(
            text=text,
            markdown=self._strip_rag_markers(text),
            page_count=None,
            meta={
                "parser": "openpyxl",
                "sheets": sheet_meta,
                "image_count": image_count,
                "ocr_enabled": self.settings.ocr_enabled,
                "ocr_engine": "paddleocr",
                "image_ocr_count": len(image_ocr_results),
                "image_ocr_results": image_ocr_results,
                "image_ocr_errors": image_ocr_errors,
                "image_ocr_required": len(image_ocr_results) < image_count,
            },
        )

    def _parse_xls_with_fallback(self, data: bytes) -> ParsedDocument:
        converter = self._office_converter()
        converted = self._convert_legacy_office(data, ".xls", "xlsx") if converter else None
        if converted:
            parsed = self._parse_xlsx(converted)
            parsed.meta = {
                **parsed.meta,
                "parser": "libreoffice+openpyxl",
                "converted_from": ".xls",
            }
            return parsed
        try:
            return self._parse_xls(data)
        except Exception as direct_error:
            if converter:
                raise RuntimeError(
                    "旧版 Excel .xls 直接解析和 LibreOffice 转换均失败，文件可能损坏、加密或格式不兼容"
                ) from direct_error
            return self._needs_conversion_document(
                ".xls",
                "旧版 Excel .xls 需要安装 xlrd，或通过 LibreOffice / Microsoft Excel 转换为 .xlsx。",
                extra_meta={"direct_parse_error": str(direct_error)},
            )

    def _parse_xls(self, data: bytes) -> ParsedDocument:
        try:
            import xlrd
        except Exception as exc:
            raise RuntimeError("解析 .xls 需要安装 xlrd") from exc

        try:
            workbook = xlrd.open_workbook(file_contents=data, formatting_info=True)
        except Exception:
            workbook = xlrd.open_workbook(file_contents=data)

        sections: list[str] = []
        sheet_meta: list[dict] = []
        for sheet_index in range(workbook.nsheets):
            sheet = workbook.sheet_by_index(sheet_index)
            rows, merged_ranges = self._xlrd_rows(workbook, sheet)
            sheet_text, metadata = self._excel_sheet_markdown(
                sheet.name,
                rows,
                merged_ranges,
                parser="xlrd",
                sheet_index=sheet_index + 1,
            )
            sections.append(sheet_text)
            sheet_meta.append(metadata)
        text = "\n\n".join(sections)
        return ParsedDocument(
            text=text,
            markdown=self._strip_rag_markers(text),
            page_count=None,
            meta={"parser": "xlrd", "sheets": sheet_meta},
        )

    def _parse_image(self, file_name: str, data: bytes) -> ParsedDocument:
        suffix = Path(file_name).suffix.lower() or ".png"
        ocr_result = self.ocr_service.ocr_image_bytes(data, suffix) if self._should_ocr(0) else None
        if ocr_result and ocr_result.text:
            metadata = {
                "source": "image",
                "file_name": file_name,
                **self._ocr_summary(ocr_result),
            }
            text = self._rag_block("image_ocr", metadata, ocr_result.text)
            return ParsedDocument(
                text=text,
                markdown=self._strip_rag_markers(text),
                page_count=None,
                meta={
                    "parser": "image-ocr",
                    "ocr_enabled": self.settings.ocr_enabled,
                    "ocr_engine": "paddleocr",
                    "ocr_required": False,
                    "ocr_result": self._ocr_summary(ocr_result),
                },
            )

        error = ocr_result.error if ocr_result else self._ocr_skip_reason()
        text = f"[图片文件未提取到可用文字，PaddleOCR 未返回可用文字：{error}]"
        return ParsedDocument(
            text=text,
            markdown=text,
            page_count=None,
            meta={
                "parser": "image-ocr",
                "ocr_enabled": self.settings.ocr_enabled,
                "ocr_engine": "paddleocr",
                "ocr_required": True,
                "ocr_errors": [{"image": 1, "error": error}],
                "skip_chunking": True,
            },
        )

    def _parse_container(self, file_name: str, data: bytes) -> ParsedDocument:
        ext = Path(file_name).suffix.lower()
        entries: list[dict] = []
        unsafe_entries: list[str] = []
        if ext == ".zip":
            import zipfile

            try:
                with zipfile.ZipFile(self._bytes_io(data)) as archive:
                    for info in archive.infolist()[:200]:
                        entry_name = self._decode_zip_name(info)
                        if self._unsafe_archive_name(entry_name):
                            unsafe_entries.append(entry_name)
                        entries.append(
                            {
                                "name": entry_name,
                                "size": info.file_size,
                                "compressed_size": info.compress_size,
                                "is_dir": info.is_dir(),
                            }
                        )
            except Exception as exc:
                entries.append({"error": str(exc)})

        lines = [
            "压缩包文件，默认不直接混入知识库正文。",
            "需要按安全策略解压后，再导入内部文档。",
        ]
        if entries:
            lines.append("内部文件预览：")
            for entry in entries[:30]:
                if "name" in entry:
                    lines.append(f"- {entry['name']} ({entry.get('size', 0)} bytes)")
                else:
                    lines.append(f"- 读取失败：{entry.get('error')}")
        text = "\n".join(lines)
        return ParsedDocument(
            text=text,
            markdown=text,
            page_count=None,
            meta={
                "parser": "archive-container",
                "container": True,
                "needs_extraction": True,
                "skip_chunking": True,
                "entry_count": len(entries),
                "entries_preview": entries,
                "unsafe_entries": unsafe_entries,
            },
        )

    def _docx_table_rows(self, table) -> list[list[str]]:
        rows: list[list[str]] = []
        for row in table.rows:
            values = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            while values and not values[-1]:
                values.pop()
            if any(values):
                rows.append(values)
        return rows

    def _openpyxl_rows(self, sheet) -> list[tuple[int, list[str]]]:
        merged_values: dict[tuple[int, int], str] = {}
        for merged_range in sheet.merged_cells.ranges:
            value = sheet.cell(merged_range.min_row, merged_range.min_col).value
            normalized = self._cell_text(value)
            if merged_range.min_col == merged_range.max_col:
                for row in range(merged_range.min_row, merged_range.max_row + 1):
                    merged_values[(row, merged_range.min_col)] = normalized
            elif merged_range.min_row != merged_range.max_row:
                for row in range(merged_range.min_row + 1, merged_range.max_row + 1):
                    merged_values[(row, merged_range.min_col)] = normalized

        rows: list[tuple[int, list[str]]] = []
        for row_index in range(1, sheet.max_row + 1):
            values = [
                merged_values.get(
                    (row_index, col_index),
                    self._cell_text(sheet.cell(row_index, col_index).value),
                )
                for col_index in range(1, sheet.max_column + 1)
            ]
            while values and not values[-1]:
                values.pop()
            if any(values):
                rows.append((row_index, values))
        return rows

    def _openpyxl_image_ocr_blocks(
        self,
        sheet,
        sheet_index: int,
        start_image_index: int,
    ) -> tuple[list[str], list[dict], list[dict], int, int]:
        images = list(getattr(sheet, "_images", []) or [])
        blocks: list[str] = []
        results: list[dict] = []
        errors: list[dict] = []
        current_index = start_image_index
        for sheet_image_index, image in enumerate(images, start=1):
            current_index += 1
            metadata = {
                "source": "excel",
                "sheet_index": sheet_index,
                "sheet_name": sheet.title,
                "image_index": current_index,
                "sheet_image_index": sheet_image_index,
                **self._openpyxl_image_anchor(image),
            }
            if not self._should_ocr(current_index - 1):
                errors.append({**metadata, "error": self._ocr_skip_reason()})
                continue
            try:
                image_bytes = self._openpyxl_image_bytes(image)
            except Exception as exc:
                errors.append({**metadata, "error": str(exc)})
                continue
            ocr_result = self.ocr_service.ocr_image_bytes(image_bytes, self._openpyxl_image_suffix(image))
            if ocr_result.text:
                summary = self._ocr_summary(ocr_result)
                block_meta = {**metadata, **summary}
                blocks.append(
                    self._rag_block(
                        "image_ocr",
                        block_meta,
                        f"### 工作表 {sheet.title} 图片 {sheet_image_index} OCR\n{ocr_result.text}",
                    )
                )
                results.append(block_meta)
            else:
                errors.append({**metadata, "error": ocr_result.error or "PaddleOCR 未返回可用文字"})
        return blocks, results, errors, len(images), current_index

    def _openpyxl_image_bytes(self, image) -> bytes:
        data_method = getattr(image, "_data", None)
        if callable(data_method):
            data = data_method()
            if data:
                return data
        ref = getattr(image, "ref", None)
        if hasattr(ref, "read"):
            data = ref.read()
            if data:
                return data
        raise ValueError("Excel 图片数据为空或无法读取")

    def _openpyxl_image_suffix(self, image) -> str:
        path = str(getattr(image, "path", "") or "")
        suffix = Path(path).suffix.lower()
        if suffix:
            return suffix
        image_format = str(getattr(image, "format", "") or "").lower()
        if image_format:
            return f".{image_format.lstrip('.')}"
        return ".png"

    def _openpyxl_image_anchor(self, image) -> dict:
        marker = getattr(getattr(image, "anchor", None), "_from", None)
        row = getattr(marker, "row", None)
        col = getattr(marker, "col", None)
        if isinstance(row, int) and isinstance(col, int):
            row_no = row + 1
            col_no = col + 1
            return {
                "anchor_row": row_no,
                "anchor_col": col_no,
                "anchor_cell": f"{self._excel_col_name(col_no)}{row_no}",
            }
        return {}

    def _xlrd_rows(self, workbook, sheet) -> tuple[list[tuple[int, list[str]]], list[str]]:
        import xlrd

        merged_values: dict[tuple[int, int], str] = {}
        merged_ranges: list[str] = []
        for row_low, row_high, col_low, col_high in getattr(sheet, "merged_cells", []):
            value = self._xlrd_cell_text(workbook, sheet, row_low, col_low)
            merged_ranges.append(
                f"{self._excel_col_name(col_low + 1)}{row_low + 1}:"
                f"{self._excel_col_name(col_high)}{row_high}"
            )
            if col_high - col_low == 1:
                for row_index in range(row_low, row_high):
                    merged_values[(row_index, col_low)] = value
            elif row_high - row_low > 1:
                for row_index in range(row_low + 1, row_high):
                    merged_values[(row_index, col_low)] = value

        rows: list[tuple[int, list[str]]] = []
        for row_index in range(sheet.nrows):
            values = [
                merged_values.get(
                    (row_index, col_index),
                    self._xlrd_cell_text(workbook, sheet, row_index, col_index),
                )
                for col_index in range(sheet.ncols)
            ]
            while values and not values[-1]:
                values.pop()
            if any(values):
                rows.append((row_index + 1, values))
        return rows, merged_ranges

    def _xlrd_cell_text(self, workbook, sheet, row_index: int, col_index: int) -> str:
        import xlrd

        cell = sheet.cell(row_index, col_index)
        if cell.ctype == xlrd.XL_CELL_EMPTY:
            return ""
        if cell.ctype == xlrd.XL_CELL_DATE:
            try:
                value = xlrd.xldate_as_datetime(cell.value, workbook.datemode)
                if value.hour == 0 and value.minute == 0 and value.second == 0:
                    return value.strftime("%Y-%m-%d")
                return value.strftime("%Y-%m-%d %H:%M:%S")
            except Exception:
                return str(cell.value).strip()
        if cell.ctype == xlrd.XL_CELL_NUMBER and float(cell.value).is_integer():
            return str(int(cell.value))
        return str(cell.value).strip()

    def _excel_sheet_markdown(
        self,
        sheet_name: str,
        rows: list[tuple[int, list[str]]],
        merged_ranges: list[str],
        parser: str,
        sheet_index: int,
    ) -> tuple[str, dict]:
        if not rows:
            metadata = {
                "source": "excel",
                "parser": parser,
                "sheet_index": sheet_index,
                "sheet_name": sheet_name,
                "row_count": 0,
                "column_count": 0,
                "merged_ranges": merged_ranges,
            }
            return f"## 工作表：{sheet_name}\n[空工作表]", metadata

        header_pos = self._detect_header_position(rows)
        table_rows = rows[header_pos:] if header_pos is not None else rows
        title_rows = rows[:header_pos] if header_pos else []
        header_row_no = table_rows[0][0]
        row_start = table_rows[0][0]
        row_end = table_rows[-1][0]
        header_values = table_rows[0][1]
        column_count = max(len(values) for _, values in table_rows)
        index_mode = "full"
        indexed_row_count = len(table_rows)
        if self._is_large_table(table_rows):
            index_mode = "sampled"
            sample_limit = max(1, int(self.settings.table_sample_rows))
            table_rows_for_index = table_rows[:sample_limit]
            indexed_row_count = len(table_rows_for_index)
        else:
            table_rows_for_index = table_rows
        metadata = {
            "source": "excel",
            "parser": parser,
            "sheet_index": sheet_index,
            "sheet_name": sheet_name,
            "row_start": row_start,
            "row_end": row_end,
            "header_row": header_row_no,
            "row_count": len(table_rows),
            "column_count": column_count,
            "fields": [value for value in header_values if value],
            "index_mode": index_mode,
            "indexed_row_count": indexed_row_count,
            "full_index_limited": index_mode != "full",
            "merged_ranges": merged_ranges,
        }
        lines = [f"## 工作表：{sheet_name}"]
        for row_no, values in title_rows:
            title = " ".join(value for value in values if value)
            if title:
                lines.append(f"说明行 {row_no}: {title}")
        if index_mode != "full":
            lines.extend(self._large_table_summary_lines(sheet_name, table_rows, table_rows_for_index, metadata))
        table = self._markdown_table([values for _, values in table_rows_for_index])
        lines.append(
            self._rag_block(
                "table",
                metadata,
                f"### {sheet_name}\n{table}",
            )
        )
        return "\n".join(lines), metadata

    def _is_large_table(self, table_rows: list[tuple[int, list[str]]]) -> bool:
        if not table_rows:
            return False
        max_rows = max(1, int(self.settings.table_full_index_max_rows))
        if len(table_rows) > max_rows:
            return True
        text_chars = sum(sum(len(value) for value in values) for _, values in table_rows)
        return text_chars > max(1, int(self.settings.max_embedding_chars_per_document))

    def _large_table_summary_lines(
        self,
        sheet_name: str,
        table_rows: list[tuple[int, list[str]]],
        indexed_rows: list[tuple[int, list[str]]],
        metadata: dict,
    ) -> list[str]:
        fields = metadata.get("fields") or []
        row_count = metadata.get("row_count", len(table_rows))
        column_count = metadata.get("column_count", 0)
        sample_rows = max(0, len(indexed_rows) - 1)
        lines = [
            "[大表受限索引]",
            f"工作表：{sheet_name}",
            f"总行数：{row_count}",
            f"总列数：{column_count}",
            f"已索引样例行数：{sample_rows}",
        ]
        if fields:
            lines.append("字段：" + "、".join(str(field) for field in fields))
        lines.append("说明：该工作表超过全文索引预算，默认只索引字段摘要和前若干行样例；如需逐行检索，可对该文档执行深度索引。")
        return lines

    def _detect_header_position(self, rows: list[tuple[int, list[str]]]) -> int | None:
        if not rows:
            return None
        candidates = rows[: min(len(rows), 15)]
        max_non_empty = max(self._non_empty_count(values) for _, values in candidates)
        if max_non_empty <= 1:
            return 0

        best_index = 0
        best_score = -1.0
        for index, (_, values) in enumerate(candidates):
            non_empty = self._non_empty_count(values)
            if non_empty < 2:
                continue
            text_cells = sum(1 for value in values if value and not self._looks_numeric(value))
            numeric_cells = sum(1 for value in values if value and self._looks_numeric(value))
            score = non_empty * 2 + text_cells - numeric_cells - index * 0.15
            if non_empty >= max(2, int(max_non_empty * 0.75)):
                score += 3
            if score > best_score:
                best_score = score
                best_index = index
        return best_index

    def _markdown_table(self, rows: list[list[str]]) -> str:
        if not rows:
            return ""
        width = max(len(row) for row in rows)
        normalized = [row + [""] * (width - len(row)) for row in rows]
        header = normalized[0]
        lines = [
            "| " + " | ".join(self._escape_table_cell(value) for value in header) + " |",
            "| " + " | ".join("---" for _ in range(width)) + " |",
        ]
        for row in normalized[1:]:
            lines.append("| " + " | ".join(self._escape_table_cell(value) for value in row) + " |")
        return "\n".join(lines)

    def _rag_block(self, block_type: str, metadata: dict, content: str) -> str:
        payload = {"type": block_type, **metadata}
        raw_meta = json.dumps(payload, ensure_ascii=False, separators=(",", ":"))
        return f"<!-- RAG_BLOCK {raw_meta} -->\n{content.strip()}\n<!-- RAG_BLOCK_END -->"

    def _strip_rag_markers(self, text: str) -> str:
        text = re.sub(r"<!--\s*RAG_BLOCK\s+\{.*?\}\s*-->\s*", "", text, flags=re.DOTALL)
        return re.sub(r"\s*<!--\s*RAG_BLOCK_END\s*-->", "", text)

    def _needs_conversion_document(
        self,
        ext: str,
        reason: str,
        extra_meta: dict | None = None,
    ) -> ParsedDocument:
        text = f"[{ext} 文件暂未完成正文解析]\n{reason}"
        meta = {
            "parser": "legacy-office-placeholder",
            "needs_conversion": True,
            "skip_chunking": True,
            "source_extension": ext,
            "conversion_attempted": True,
            "conversion_available": bool(self._office_converter()),
        }
        if extra_meta:
            meta.update(extra_meta)
        return ParsedDocument(text=text, markdown=text, page_count=None, meta=meta)

    def convert_legacy_office_file(self, file_name: str, data: bytes) -> tuple[str, bytes, dict]:
        ext = Path(file_name).suffix.lower()
        targets = {
            ".doc": ".docx",
            ".xls": ".xlsx",
        }
        if ext not in targets:
            raise ValueError(f"当前文件不是可转换的旧 Office 格式: {ext}")
        target_ext = targets[ext]
        converted = self._convert_legacy_office(data, ext, target_ext.lstrip("."))
        if not converted:
            converter = self._office_converter()
            if not converter:
                raise RuntimeError("未找到 LibreOffice/soffice，无法转换旧 Office 文件")
            raise RuntimeError("LibreOffice 转换失败，文件可能损坏、加密或格式不兼容")
        converted_name = f"{Path(file_name).stem}{target_ext}"
        return (
            converted_name,
            converted,
            {
                "converted_from": ext,
                "converted_to": target_ext,
                "converter": self._office_converter(),
            },
        )

    def _convert_legacy_office(self, data: bytes, source_ext: str, target_ext: str) -> bytes | None:
        executable = self._office_converter()
        if not executable:
            return None
        with TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            source = temp_path / f"input{source_ext}"
            profile = temp_path / "lo-profile"
            profile.mkdir(parents=True, exist_ok=True)
            source.write_bytes(data)
            command = [
                executable,
                f"-env:UserInstallation={profile.as_uri()}",
                "--headless",
                "--convert-to",
                target_ext,
                "--outdir",
                str(temp_path),
                str(source),
            ]
            try:
                completed = subprocess.run(
                    command,
                    cwd=temp_dir,
                    capture_output=True,
                    timeout=90,
                    check=False,
                )
            except Exception:
                return None
            if completed.returncode != 0:
                return None
            candidates = sorted(temp_path.glob(f"*.{target_ext}"))
            if not candidates:
                return None
            return candidates[0].read_bytes()

    def _office_converter(self) -> str | None:
        for name in ("soffice", "libreoffice"):
            found = shutil.which(name)
            if found:
                return found
        for path in (
            r"C:\Program Files\LibreOffice\program\soffice.exe",
            r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
        ):
            if Path(path).exists():
                return path
        return None

    def _pdf_text_needs_ocr(self, page: Any, text: str) -> bool:
        compact = re.sub(r"\s+", "", text)
        if not compact:
            return True
        replacement_ratio = compact.count("�") / max(len(compact), 1)
        if replacement_ratio > 0.08:
            return True
        image_count = len(page.get_images(full=True))
        if image_count and len(compact) < 40:
            return True
        return False

    def _prefer_ocr_text(self, extracted_text: str, ocr_text: str) -> bool:
        if not extracted_text.strip():
            return True
        compact_extracted = re.sub(r"\s+", "", extracted_text)
        compact_ocr = re.sub(r"\s+", "", ocr_text)
        if compact_extracted.count("�") > compact_ocr.count("�"):
            return True
        return len(compact_ocr) >= max(20, int(len(compact_extracted) * 1.25))

    def _cell_text(self, value: Any) -> str:
        if value is None:
            return ""
        return str(value).strip()

    def _escape_table_cell(self, value: str) -> str:
        return value.replace("|", "\\|").replace("\n", " ").strip()

    def _non_empty_count(self, values: list[str]) -> int:
        return sum(1 for value in values if value.strip())

    def _looks_numeric(self, value: str) -> bool:
        clean = value.strip().replace(",", "").replace("%", "")
        if not clean:
            return False
        try:
            float(clean)
            return True
        except ValueError:
            return False

    def _excel_col_name(self, col_index: int) -> str:
        name = ""
        while col_index:
            col_index, remainder = divmod(col_index - 1, 26)
            name = chr(65 + remainder) + name
        return name

    def _docx_heading_level(self, paragraph) -> int:
        style_name = (getattr(paragraph.style, "name", "") or "").lower()
        match = re.search(r"heading\s*(\d+)", style_name)
        if match:
            return int(match.group(1))
        match = re.search(r"标题\s*(\d+)", style_name)
        if match:
            return int(match.group(1))
        return 0

    def _unsafe_archive_name(self, name: str) -> bool:
        normalized = name.replace("\\", "/")
        return normalized.startswith("/") or ".." in normalized.split("/")

    def _decode_zip_name(self, info: zipfile.ZipInfo) -> str:
        name = info.filename
        if info.flag_bits & 0x800:
            return name
        try:
            raw = name.encode("cp437")
        except UnicodeEncodeError:
            return name
        for encoding in ("utf-8", "gbk", "gb2312", "big5"):
            try:
                decoded = raw.decode(encoding)
            except UnicodeDecodeError:
                continue
            if decoded == name:
                return decoded
            if self._contains_cjk(decoded) or self._looks_mojibake(name):
                return decoded
        return name

    def _contains_cjk(self, value: str) -> bool:
        return any("\u4e00" <= char <= "\u9fff" for char in value)

    def _looks_mojibake(self, value: str) -> bool:
        return any(char in value for char in ("�", "╬", "─", "╓", "╨", "▒", "▓", "│"))

    def _bytes_io(self, data: bytes):
        from io import BytesIO

        return BytesIO(data)

    def _should_ocr(self, attempted_count: int) -> bool:
        return bool(self.settings.ocr_enabled) and attempted_count < self.settings.ocr_max_pages

    def _ocr_skip_reason(self) -> str:
        if not self.settings.ocr_enabled:
            return "OCR is disabled"
        return f"OCR attempt limit reached: {self.settings.ocr_max_pages}"

    def _ocr_pdf_page(self, page) -> OcrResult:
        dpi = max(96, int(self.settings.ocr_pdf_dpi))
        zoom = dpi / 72
        if dpi != 72:
            import fitz

            pixmap = page.get_pixmap(matrix=fitz.Matrix(zoom, zoom), alpha=False)
        else:
            pixmap = page.get_pixmap(alpha=False)
        return self.ocr_service.ocr_image_bytes(pixmap.tobytes("png"), ".png")

    def _ocr_summary(self, result: OcrResult) -> dict:
        if not result.lines:
            return {"line_count": 0, "avg_confidence": 0}
        avg_confidence = sum(line.confidence for line in result.lines) / len(result.lines)
        return {
            "line_count": len(result.lines),
            "avg_confidence": round(avg_confidence, 4),
        }

    def _image_suffix(self, content_type: str) -> str:
        suffixes = {
            "image/png": ".png",
            "image/jpeg": ".jpg",
            "image/jpg": ".jpg",
            "image/bmp": ".bmp",
            "image/tiff": ".tiff",
            "image/gif": ".gif",
        }
        return suffixes.get(content_type.lower(), ".png")

    def quality_score(self, parsed: ParsedDocument) -> float:
        text = parsed.text.strip()
        if not text:
            return 0.0
        if parsed.meta.get("needs_conversion") or parsed.meta.get("needs_extraction"):
            return 20.0
        if parsed.meta.get("skip_chunking"):
            return 25.0
        score = 70.0
        if len(text) > 2000:
            score += 15
        if parsed.page_count:
            score += 5
        if parsed.meta.get("ocr_required") or parsed.meta.get("image_ocr_required"):
            score -= 15
        if parsed.meta.get("table_count") or parsed.meta.get("sheets"):
            score += 5
        if "�" not in text:
            score += 10
        return min(score, 100.0)
